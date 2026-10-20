"""
UTS BUSINESS INTELLIGENCE - ANALISIS TRANSAKSI TOKO (EDA, PREPROCESSING, APRIORI, K-MEANS, DASHBOARD BI)
Universitas Timor (UNIMOR)

Alur: LOGIN -> Beranda -> (upload Excel) -> 1 Business Understanding -> 2 EDA -> 3 Preprocessing & Feature Engineering
      -> 4 Apriori -> 5 K-Means (RFM) -> 6 Dashboard BI -> 7 Rekomendasi Keputusan

Jalankan:
    pip install streamlit pandas numpy plotly openpyxl pillow
    streamlit run app_uts_bi_tokotoko.py

Login : username = helga | password = helga123  (ubah di bagian AUTH_USER / AUTH_PASS)

Data: "P8__UTS_Business_Intelligence.xlsx" (sheet Data_Raw, Data_Produk, Data_Pelanggan).
 - Jika file ini diletakkan di folder yang SAMA dengan program, data dibaca otomatis.
 - Bisa juga di-upload lewat halaman Analisis.

Aset opsional (folder yang sama): bg_unimor.jpg (background), logo_unimor.png (logo), foto.jpg (foto profil).
Semua algoritma (Apriori, K-Means, Silhouette) ditulis dengan numpy/pandas sehingga tanpa scikit-learn / mlxtend.
"""
import base64
import hmac
import io
import itertools
from pathlib import Path

import numpy as np
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import streamlit as st
from plotly.subplots import make_subplots

st.set_page_config(page_title="UTS BI | Analisis Transaksi Toko - UNIMOR", page_icon="📊",
                   layout="wide", initial_sidebar_state="collapsed")

# ----------------------------------------------------------------------------
# IDENTITAS (ubah di sini)
# ----------------------------------------------------------------------------
nama, npm, kelas, semester = "Helga Lydia S. Funan", "51230058", "Kelas B", "Semester 7"
matkul = "Business Intelligence"
UNIMOR_NAMA = "Universitas Timor (UNIMOR)"
UNIMOR_TAGLINE = "Perguruan Tinggi Negeri di Kefamenanu, Kabupaten Timor Tengah Utara, Nusa Tenggara Timur"
UNIMOR_ALAMAT = "Jl. El Tari Km 09, Kelurahan Sasi, Kecamatan Kefamenanu, Kab. Timor Tengah Utara, NTT 85613"
UNIMOR_WEB = "https://unimor.ac.id"
UNIMOR_FAKULTAS = [
    ("📚", "Fakultas Keguruan & Ilmu Pendidikan", "Mencetak tenaga pendidik profesional."),
    ("🏛️", "Fakultas Ilmu Sosial & Ilmu Politik", "Kajian sosial, politik, dan pemerintahan."),
    ("💼", "Fakultas Ekonomi & Bisnis", "Ekonomi, manajemen, dan bisnis."),
    ("🌾", "Fakultas Pertanian, Sains & Kesehatan", "Pertanian, sains, dan kesehatan."),
]

# ----------------------------------------------------------------------------
# AKUN LOGIN (ubah di sini)
# ----------------------------------------------------------------------------
AUTH_USER = "helga"
AUTH_PASS = "helga123"

# ----------------------------------------------------------------------------
# WARNA  (tema UNIMOR + palet khusus agar tiap grafik berbeda warna)
# ----------------------------------------------------------------------------
NAVY, BLUE, TEAL, GOLD, LIGHT = "#0B2A4A", "#14578F", "#1D9AA6", "#F2B705", "#F4F7FB"
C = dict(pink="#E0457B", purple="#7C4DFF", orange="#F57C00", green="#2E933C", cyan="#00A6B8",
         indigo="#3F51B5", red="#E4572E", lime="#8BC34A", brown="#8D6E63", magenta="#C2185B",
         sky="#29B6F6", amber="#FFB300", teal="#00897B", slate="#607D8B")
CAT_COLORS = ["#14578F", "#E0457B", "#2E933C", "#F57C00", "#7C4DFF", "#00A6B8", "#FFB300", "#8D6E63", "#C2185B"]
SEG_PALETTE = ["#2E933C", "#14578F", "#F57C00", "#E4572E", "#7C4DFF", "#00A6B8", "#C2185B", "#8D6E63"]
SEG_ICON = {"Loyal": "👑", "Potential": "🚀", "New": "🌱", "Churn": "⚠️"}

APP_DIR = Path(__file__).parent
DEFAULT_FILE_NAMES = ["P8__UTS_Business_Intelligence.xlsx", "P8_UTS_Business_Intelligence.xlsx", "data_uts.xlsx"]
UNKNOWN = "Tidak Diketahui"
PAY_CANON = {"tunai": "Tunai", "debit": "Debit", "transfer": "Transfer", "qris": "QRIS"}
MAX_QTY_WAJAR = 20  # batas atas kuantitas per baris (data valid maksimum 6; 999 = placeholder)
BULAN = ["Jan", "Feb", "Mar", "Apr", "Mei", "Jun", "Jul", "Agu", "Sep", "Okt", "Nov", "Des"]

CSS = f"""
<style>
:root, html, body, .stApp {{ color-scheme: light; }}
.stApp {{ color:#1D2B3A; font-family:"Segoe UI", "Inter", Arial, sans-serif; }}
.block-container {{ padding-top: 1rem; max-width: 1350px; }}
section[data-testid="stSidebar"], div[data-testid="stSidebarCollapsedControl"], button[data-testid="stExpandSidebarButton"] {{ display:none !important; }}
header[data-testid='stHeader'] {{ background: transparent; }}
h1, h2, h3, h4, h5 {{ color:{NAVY}; letter-spacing:-.2px; }}
h3 {{ font-weight:800; }}
h4 {{ font-weight:700; margin-top:.6rem; }}

/* ---------- HERO (halaman analisis) ---------- */
.hero {{ background: linear-gradient(120deg, #EAF3FF 0%, #E8F6F8 100%); border:1px solid #D6E4F2; border-radius: 22px; padding: 24px 30px;
    color:{NAVY}; display:flex; align-items:center; gap:26px; flex-wrap:wrap; border-bottom:6px solid {GOLD};
    box-shadow:0 8px 24px rgba(20,87,143,.10); }}
.hero img.pf, .hero .avatar {{ width:120px; height:120px; border-radius:50%; object-fit:cover; border:5px solid {TEAL}; background:#fff; flex-shrink:0; }}
.hero .avatar {{ display:flex; align-items:center; justify-content:center; font-size:42px; font-weight:800; color:{NAVY}; }}
.hero img.lg {{ width:88px; height:88px; object-fit:contain; background:#fff; border-radius:50%; padding:6px; border:3px solid {GOLD}; flex-shrink:0; }}
.badge {{ background:{GOLD}; color:{NAVY}; font-weight:800; letter-spacing:1px; padding:4px 14px; border-radius:999px; font-size:12px; display:inline-block; }}
.hero h1 {{ margin:8px 0 4px 0; font-size:28px; color:{NAVY}; line-height:1.2; }}
.hero p.sub {{ margin:0 0 12px 0; color:#4A5B6E; font-size:14.5px; }}
.chips span {{ display:inline-block; background:#fff; color:{NAVY}; border:1px solid #D6E4F2; padding:5px 13px; border-radius:12px; margin:0 8px 6px 0; font-size:13.5px; }}
.chips b {{ color:{BLUE}; }}

/* ---------- KPI ---------- */
.kpi {{ background:#fff; border-radius:16px; padding:14px 16px; border:1px solid #E3E9F1; border-left:6px solid {TEAL};
    box-shadow:0 3px 12px rgba(20,87,143,.07); height:100%; transition:transform .15s ease, box-shadow .15s ease; }}
.kpi:hover {{ transform:translateY(-2px); box-shadow:0 8px 20px rgba(20,87,143,.14); }}
.kpi .lbl {{ font-size:11.5px; color:#5B6B7F; font-weight:700; text-transform:uppercase; letter-spacing:.5px; }}
.kpi .val {{ font-size:24px; font-weight:800; color:{NAVY}; line-height:1.25; }}
.kpi .sub {{ font-size:12px; color:#6B7B8E; }}
.kpi.gold {{ border-left-color:{GOLD}; }} .kpi.blue {{ border-left-color:{BLUE}; }} .kpi.red {{ border-left-color:{C['red']}; }}
.kpi.green {{ border-left-color:{C['green']}; }} .kpi.pink {{ border-left-color:{C['pink']}; }} .kpi.purple {{ border-left-color:{C['purple']}; }}
.kpi.orange {{ border-left-color:{C['orange']}; }}

/* ---------- KOTAK INFO ---------- */
.note {{ background:#FFF8E1; border-left:5px solid {GOLD}; padding:12px 16px; border-radius:12px; color:#4A3B00; line-height:1.55; }}
.ok {{ background:#E8F6EC; border-left:5px solid {C['green']}; padding:12px 16px; border-radius:12px; color:#14471E; line-height:1.55; }}
.info {{ background:#E8F1FA; border-left:5px solid {BLUE}; padding:12px 16px; border-radius:12px; color:{NAVY}; line-height:1.55; }}
.warn {{ background:#FDECEA; border-left:5px solid {C['red']}; padding:12px 16px; border-radius:12px; color:#5A1A0E; line-height:1.55; }}
.tip {{ background:linear-gradient(90deg,#F1F8FF,#F0FAF9); border:1px dashed #9CC4E4; border-radius:14px; padding:10px 16px; color:{NAVY}; font-size:14px; }}
.tip b {{ color:{BLUE}; }}

/* ---------- TAB / TOMBOL / INPUT ---------- */
div[data-testid="stTabs"] [role="tablist"] {{ gap:6px; flex-wrap:wrap; border-bottom:2px solid #E3E9F1; }}
div[data-testid="stTabs"] button[role="tab"] {{ font-weight:700; color:#4A5B6E; background:#F1F6FB; border-radius:12px 12px 0 0; padding:8px 14px; }}
div[data-testid="stTabs"] button[role="tab"]:hover {{ background:#E3EEF9; color:{NAVY}; }}
div[data-testid="stTabs"] button[role="tab"][aria-selected="true"] {{ background:#fff; color:{BLUE}; box-shadow:inset 0 -3px 0 {GOLD}; }}
div.stButton > button, div.stDownloadButton > button, div[data-testid="stFormSubmitButton"] > button {{ border-radius:12px; font-weight:700; border:1px solid #C9D9EA; background:#fff; color:{NAVY};
    transition:all .15s ease; }}
div.stButton > button:hover, div.stDownloadButton > button:hover {{ border-color:{BLUE}; color:{BLUE}; background:#F3F8FD; transform:translateY(-1px); }}
div.stButton > button[kind="primary"], div[data-testid="stFormSubmitButton"] > button[kind="primaryFormSubmit"], div[data-testid="stFormSubmitButton"] > button[kind="primary"] {{
    background:linear-gradient(90deg,{BLUE},{TEAL}); color:#fff; border:none; box-shadow:0 4px 12px rgba(20,87,143,.25); }}
div.stButton > button[kind="primary"]:hover {{ color:#fff; filter:brightness(1.06); }}
div[data-testid="stExpander"] {{ border:1px solid #E3E9F1; border-radius:14px; background:#FBFDFF; }}
div[data-testid="stExpander"] summary p {{ font-weight:700; color:{NAVY}; }}
div[data-testid="stVerticalBlockBorderWrapper"] {{ border-radius:16px; }}
[data-testid="stFileUploader"] section {{ border:2px dashed {TEAL}; background:{LIGHT}; border-radius:14px; }}

/* ---------- ALUR / LANGKAH ---------- */
.flow {{ display:flex; gap:8px; flex-wrap:wrap; margin:6px 0 14px 0; }}
.step {{ flex:1; min-width:120px; background:#fff; border:1px solid #D6E4F2; border-top:4px solid {GOLD}; color:{NAVY}; padding:12px 14px; border-radius:14px; text-align:center;
    box-shadow:0 3px 10px rgba(20,87,143,.06); }}
.step .n {{ background:{BLUE}; color:#fff; font-weight:800; width:26px; height:26px; border-radius:50%; display:inline-flex; align-items:center; justify-content:center; margin-bottom:5px; }}
.step .t {{ font-weight:700; font-size:14px; color:{NAVY}; }} .step .d {{ font-size:11.5px; color:#5B6B7F; }}

/* ---------- KARTU SEGMEN ---------- */
.seg {{ --sc:{TEAL}; background:#fff; border:1px solid #E3E9F1; border-top:6px solid var(--sc); border-radius:18px; padding:16px 18px; color:#1D2B3A; height:100%;
    box-shadow:0 5px 16px rgba(20,87,143,.09); }}
.seg h4 {{ margin:0 0 2px 0; color:var(--sc); font-size:19px; }}
.seg .tag {{ font-size:11.5px; color:#5B6B7F; margin-bottom:8px; letter-spacing:.5px; text-transform:uppercase; font-weight:700; }}
.seg .big {{ font-size:28px; font-weight:800; line-height:1.1; color:{NAVY}; }}
.seg .mini {{ font-size:12.5px; color:#4A5B6E; margin:2px 0 8px 0; }}
.seg ul {{ margin:6px 0 0 0; padding-left:18px; font-size:13px; line-height:1.55; color:#33475B; }}
.seg .rfm {{ display:flex; gap:8px; margin:8px 0; }}
.seg .rfm div {{ flex:1; background:#F1F6FB; border-radius:10px; padding:6px 8px; text-align:center; }}
.seg .rfm b {{ display:block; font-size:15px; color:{NAVY}; }} .seg .rfm span {{ font-size:10.5px; color:#5B6B7F; }}

.brandbar {{ font-weight:700; padding-top:8px; font-size:15px; color:{NAVY}; }}

/* ---------- BERANDA ---------- */
.uhero {{ background:rgba(255,255,255,.94); border:1px solid #D6E4F2; border-radius:22px; padding:28px 32px; color:{NAVY};
    display:flex; align-items:center; gap:26px; flex-wrap:wrap; border-bottom:6px solid {GOLD}; box-shadow:0 10px 28px rgba(20,87,143,.14); }}
.uhero img.logo {{ width:120px; height:120px; object-fit:contain; background:#fff; border-radius:50%; padding:8px; border:4px solid {GOLD}; flex-shrink:0; }}
.uhero .emoji {{ font-size:84px; }}
.uhero h1 {{ margin:8px 0 4px 0; font-size:36px; color:{NAVY}; line-height:1.12; }}
.uhero p {{ margin:0; color:#4A5B6E; font-size:16px; }}
.sec-title {{ color:{NAVY}; font-size:22px; font-weight:800; margin:26px 0 10px 0; padding-left:12px; border-left:6px solid {GOLD}; }}
.grid {{ display:grid; grid-template-columns:repeat(auto-fit, minmax(230px,1fr)); gap:14px; }}
.glass {{ background:#fff; border:1px solid #E3E9F1; border-radius:18px; padding:18px 20px; color:#1D2B3A; box-shadow:0 5px 16px rgba(20,87,143,.09); border-top:5px solid {TEAL}; }}
.glass.gold {{ border-top-color:{GOLD}; }} .glass.blue {{ border-top-color:{BLUE}; }} .glass.pink {{ border-top-color:{C['pink']}; }} .glass.purple {{ border-top-color:{C['purple']}; }}
.glass .ic {{ font-size:28px; }} .glass h4 {{ margin:6px 0 6px 0; color:{NAVY}; font-size:16px; }}
.glass p {{ margin:0; font-size:14px; line-height:1.55; color:#33475B; }}
.stat {{ background:#fff; border:1px solid #D6E4F2; border-radius:16px; padding:16px; text-align:center; color:{NAVY}; box-shadow:0 4px 12px rgba(20,87,143,.08); }}
.stat .v {{ font-size:26px; font-weight:800; color:{BLUE}; }} .stat .l {{ font-size:13px; color:#5B6B7F; }}
a.lnk {{ display:inline-block; background:{NAVY}; color:#fff !important; text-decoration:none; padding:8px 16px; border-radius:10px; font-weight:700; font-size:14px; margin:6px 8px 0 0; }}
.foot {{ text-align:center; color:#6B7B8E; font-size:13px; margin:28px 0 6px 0; }}
</style>
"""
st.markdown(CSS, unsafe_allow_html=True)


# ----------------------------------------------------------------------------
# FUNGSI TAMPILAN
# ----------------------------------------------------------------------------
def rupiah(x):
    return "Rp " + f"{x:,.0f}".replace(",", ".")


def rp_short(x):
    a = abs(x)
    if a >= 1e9:
        return f"Rp {x / 1e9:,.2f} M".replace(",", "X").replace(".", ",").replace("X", ".")
    if a >= 1e6:
        return f"Rp {x / 1e6:,.1f} jt".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"Rp {x / 1e3:,.0f} rb".replace(",", ".")


def angka(x):
    return f"{x:,.0f}".replace(",", ".")


def dec(x, d=2):
    return f"{x:,.{d}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def kpi(label, value, sub="", cls=""):
    return f'<div class="kpi {cls}"><div class="lbl">{label}</div><div class="val">{value}</div><div class="sub">{sub}</div></div>'


def show_chart(fig, height=None):
    fig.update_layout(template="plotly_white", margin=dict(l=10, r=10, t=55, b=10),
                      font=dict(family="Segoe UI, Arial", size=13, color="#1D2B3A"), title_font=dict(size=16, color=NAVY),
                      paper_bgcolor="rgba(0,0,0,0)")
    if height:
        fig.update_layout(height=height)
    try:
        st.plotly_chart(fig, width="stretch")
    except TypeError:
        st.plotly_chart(fig, use_container_width=True)


def show_df(df, **kw):
    try:
        st.dataframe(df, width="stretch", hide_index=True, **kw)
    except TypeError:
        st.dataframe(df, use_container_width=True, hide_index=True, **kw)


def stretch(fn, *a, **kw):
    try:
        return fn(*a, width="stretch", **kw)
    except TypeError:
        return fn(*a, use_container_width=True, **kw)


def do_rerun():
    (st.rerun if hasattr(st, "rerun") else st.experimental_rerun)()


def qa(judul, isi):
    with st.expander(f"💡 {judul}"):
        st.markdown(isi)


def sec(txt):
    st.markdown(f"#### {txt}")


def to_excel_bytes(sheets: dict) -> bytes:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl", datetime_format="yyyy-mm-dd hh:mm") as w:
        for n, f in sheets.items():
            f.to_excel(w, sheet_name=n[:31], index=False)
        for n, f in sheets.items():
            ws = w.sheets[n[:31]]
            for c in ws[1]:
                c.fill = PatternFill("solid", fgColor=NAVY.lstrip("#"))
                c.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.row_dimensions[1].height = 30
            ws.freeze_panes = "A2"
            for i, col in enumerate(f.columns, 1):
                smp = max((len(str(v)) for v in f[col].head(300)), default=0)
                ws.column_dimensions[get_column_letter(i)].width = min(max(len(str(col)) + 3, int(smp) + 2, 10), 60)
    return buf.getvalue()


# ----------------------------------------------------------------------------
# ASET (background, logo, foto)
# ----------------------------------------------------------------------------
def b64(data: bytes, mime: str):
    return f"data:{mime};base64," + base64.b64encode(data).decode()


def find_asset(names):
    for n in names:
        if (APP_DIR / n).exists():
            return APP_DIR / n
    return None


@st.cache_data(show_spinner=False)
def load_bg(path_str, mt):
    try:
        from PIL import Image
        img = Image.open(path_str).convert("RGB")
        if img.width > 1920:
            img = img.resize((1920, int(img.height * 1920 / img.width)))
        buf = io.BytesIO()
        img.save(buf, "JPEG", quality=78, optimize=True)
        return b64(buf.getvalue(), "image/jpeg")
    except Exception:
        return None


@st.cache_data(show_spinner=False)
def load_img(path_str, mt):
    p = Path(path_str)
    return b64(p.read_bytes(), "image/png" if p.suffix.lower() == ".png" else "image/jpeg")


def get_logo():
    p = find_asset(["logo_unimor.png", "logo_unimor.jpg", "logo_unimor.jpeg", "logo.png", "logo.jpg"])
    return load_img(str(p), p.stat().st_mtime) if p else None


def get_foto():
    p = find_asset(["foto.jpg", "foto.jpeg", "foto.png", "Foto.jpg", "Foto.png"])
    return load_img(str(p), p.stat().st_mtime) if p else None


def bg_css(panel: bool):
    p = find_asset(["bg_unimor.jpg", "bg_unimor.jpeg", "bg_unimor.png", "bg_unimor.webp", "background.jpg", "background.png"])
    img = load_bg(str(p), p.stat().st_mtime) if p else None
    ov = "linear-gradient(135deg, rgba(246,250,254,.94) 0%, rgba(235,244,251,.92) 55%, rgba(232,246,246,.90) 100%)"
    bg = f'{ov}, url("{img}")' if img else "linear-gradient(135deg, #F6FAFE 0%, #EAF3FA 55%, #E8F6F6 100%)"
    css = f".stApp {{ background:{bg}; background-size:cover; background-position:center; background-attachment:fixed; }}"
    if panel:  # halaman analisis: isi dalam panel putih agar grafik terbaca jelas
        css += (".block-container { background:#FFFFFF; border:1px solid #E3E9F1; border-radius:24px; margin-top:14px; margin-bottom:14px;"
                " padding:1.2rem 2rem 2rem 2rem !important; box-shadow:0 10px 34px rgba(20,87,143,.12); }")
    return f"<style>{css}</style>"


def logo_html(cls="logo"):
    lg = get_logo()
    return f'<img class="{cls}" src="{lg}" alt="Logo UNIMOR">' if lg else '<div class="emoji">🎓</div>'


# ----------------------------------------------------------------------------
# LOGIN
# ----------------------------------------------------------------------------
st.session_state.setdefault("auth", False)
st.session_state.setdefault("auth_user", "")
st.session_state.setdefault("login_error", "")
st.session_state.setdefault("page", "beranda")

LOGIN_CSS = f"""
<style>
.block-container {{ padding-top:2.2rem; max-width:1100px; }}
.lg-brand {{ text-align:center; margin-bottom:14px; }}
.lg-brand .badge {{ background:{GOLD}; color:{NAVY}; font-weight:800; letter-spacing:1px; padding:4px 14px;
    border-radius:999px; font-size:12px; display:inline-block; }}
.lg-brand h2 {{ color:{NAVY}; margin:10px 0 2px 0; font-size:28px; line-height:1.15; }}
.lg-brand p {{ margin:0; color:#4A5B6E; font-size:14px; }}
.lg-profile {{ background:#fff; border:1px solid #D6E4F2; border-bottom:none; border-radius:22px 22px 0 0; padding:22px 24px 16px 24px; text-align:center;
    border-top:6px solid {GOLD}; box-shadow:0 12px 32px rgba(20,87,143,.14); }}
.lg-profile .imgs {{ display:flex; justify-content:center; align-items:center; gap:14px; margin-bottom:10px; }}
.lg-profile img.logo {{ width:70px; height:70px; object-fit:contain; background:#fff; border-radius:50%; padding:5px; border:3px solid {GOLD}; }}
.lg-profile img.pf, .lg-profile .avatar {{ width:108px; height:108px; border-radius:50%; object-fit:cover; border:5px solid {TEAL}; background:#fff; }}
.lg-profile .avatar {{ display:inline-flex; align-items:center; justify-content:center; font-size:38px; font-weight:800; color:{NAVY}; }}
.lg-profile h3 {{ margin:6px 0 2px 0; color:{NAVY}; font-size:22px; }}
.lg-profile .sub {{ color:#5B6B7F; font-size:13.5px; margin-bottom:10px; }}
.lg-profile .chips span {{ display:inline-block; background:#EAF2FB; color:{NAVY}; border:1px solid #D3E2F3;
    padding:4px 12px; border-radius:12px; margin:0 4px 6px 0; font-size:12.5px; }}
.lg-profile .chips b {{ color:{BLUE}; }}
div[data-testid="stForm"] {{ background:#fff; border:1px solid #D6E4F2; border-top:none; border-radius:0 0 22px 22px;
    padding:6px 24px 20px 24px; box-shadow:0 12px 32px rgba(20,87,143,.14); }}
div[data-testid="stForm"] label p {{ font-weight:700; color:{NAVY}; }}
.lg-hint {{ text-align:center; margin-top:14px; color:#4A5B6E; font-size:12.5px; }}
.lg-hint code {{ background:#FFF3C4; color:#6B4E00; padding:2px 8px; border-radius:8px; font-size:12px; }}
.lg-foot {{ text-align:center; color:#6B7B8E; font-size:12px; margin-top:18px; }}
</style>
"""


def check_login(user: str, pwd: str) -> bool:
    ok_u = hmac.compare_digest(user.strip().lower().encode(), AUTH_USER.encode())
    ok_p = hmac.compare_digest(pwd.encode(), AUTH_PASS.encode())
    return ok_u and ok_p


def logout():
    st.session_state.auth = False
    st.session_state.auth_user = ""
    st.session_state.login_error = ""
    st.session_state.page = "beranda"
    do_rerun()


def render_login():
    st.markdown(bg_css(False), unsafe_allow_html=True)
    st.markdown(LOGIN_CSS, unsafe_allow_html=True)
    lg_, ft_ = get_logo(), get_foto()
    inisial_ = "".join(w[0] for w in nama.split()[:2]).upper()
    img_logo = f'<img class="logo" src="{lg_}" alt="Logo UNIMOR">' if lg_ else ""
    img_foto = f'<img class="pf" src="{ft_}" alt="Foto {nama}">' if ft_ else f'<div class="avatar">{inisial_}</div>'
    _, mid, _ = st.columns([1, 1.5, 1])
    with mid:
        st.markdown('<div class="lg-brand"><span class="badge">LOGIN • UTS BUSINESS INTELLIGENCE</span>'
                    "<h2>📊 Analisis Transaksi Toko</h2><p>Universitas Timor (UNIMOR)</p></div>", unsafe_allow_html=True)
        # profil mahasiswa
        st.markdown(f'<div class="lg-profile"><div class="imgs">{img_logo}{img_foto}</div>'
                    f"<h3>{nama}</h3><div class=\"sub\">Mahasiswa • Mata kuliah {matkul}</div>"
                    f'<div class="chips"><span>🆔 <b>NPM</b> : {npm}</span><span>🏫 <b>Kelas</b> : {kelas.replace("Kelas ", "")}</span>'
                    f"<span>📘 <b>{semester}</b></span></div></div>", unsafe_allow_html=True)
        # form login
        with st.form("form_login", clear_on_submit=False):
            user = st.text_input("👤 Username", placeholder="Masukkan username")
            pwd = st.text_input("🔒 Password", type="password", placeholder="Masukkan password")
            go_ = stretch(st.form_submit_button, "🔐 Masuk", type="primary")
        if go_:
            if not user.strip() or not pwd:
                st.session_state.login_error = "Username dan password wajib diisi."
            elif check_login(user, pwd):
                st.session_state.auth = True
                st.session_state.auth_user = user.strip().lower()
                st.session_state.login_error = ""
                st.session_state.page = "beranda"
                do_rerun()
            else:
                st.session_state.login_error = "Username atau password salah. Silakan coba lagi."
        if st.session_state.login_error:
            st.error(st.session_state.login_error, icon="🚫")
        # keterangan kecil di bawah login
        st.markdown(f'<div class="lg-hint">ℹ️ Akun demo &nbsp;•&nbsp; Username: <code>{AUTH_USER}</code> '
                    f'&nbsp;•&nbsp; Password: <code>{AUTH_PASS}</code></div>'
                    f'<div class="lg-foot">© Universitas Timor • {matkul} • {nama} ({npm})</div>', unsafe_allow_html=True)


if not st.session_state.auth:
    render_login()
    st.stop()


# ----------------------------------------------------------------------------
# TOPBAR & BERANDA
# ----------------------------------------------------------------------------
def topbar(dark):
    color = NAVY
    c1, c2, c3 = st.columns([6.2, 1.5, 1.2])
    c1.markdown(f'<div class="brandbar" style="color:{color}">🎓 UNIMOR &nbsp;•&nbsp; UTS {matkul} &nbsp;•&nbsp; {nama} ({npm}) &nbsp;•&nbsp; 👤 {st.session_state.auth_user}</div>', unsafe_allow_html=True)
    if stretch(c3.button, "🚪 Keluar", key="nav_logout"):
        logout()
    if st.session_state.page == "analisis":
        if stretch(c2.button, "🏠 Beranda", key="nav_home"):
            st.session_state.page = "beranda"
            do_rerun()
    else:
        if stretch(c2.button, "📊 Analisis", key="nav_an"):
            st.session_state.page = "analisis"
            do_rerun()


STEPS = ["Business Understanding & EDA", "Data Preprocessing", "Apriori", "K-Means (RFM)", "Dashboard BI", "Rekomendasi"]


def render_beranda():
    st.markdown(bg_css(False), unsafe_allow_html=True)
    topbar(True)
    st.markdown(f'<div class="uhero">{logo_html()}<div><span class="badge">BERANDA • IDENTITAS KAMPUS</span>'
                f"<h1>{UNIMOR_NAMA}</h1><p>{UNIMOR_TAGLINE}</p></div></div>", unsafe_allow_html=True)
    st.markdown(
        '<div class="grid" style="margin-top:16px">'
        '<div class="stat"><div class="v">2014</div><div class="l">Berstatus PTN (Perpres 119/2014)</div></div>'
        f'<div class="stat"><div class="v">{len(UNIMOR_FAKULTAS)}</div><div class="l">Fakultas</div></div>'
        '<div class="stat"><div class="v">Baik</div><div class="l">Akreditasi Institusi BAN-PT</div></div>'
        '<div class="stat"><div class="v">Kefamenanu</div><div class="l">Timor Tengah Utara, NTT</div></div></div>',
        unsafe_allow_html=True)
    st.markdown('<div class="sec-title">🏫 Profil Universitas</div>', unsafe_allow_html=True)
    cards = [("🏛️", "Status & Dasar Hukum", "Perguruan tinggi negeri yang didirikan berdasarkan Peraturan Presiden Nomor 119 Tahun 2014 tentang Pendirian Universitas Timor.", ""),
             ("📜", "Sejarah Singkat", "Berakar dari Universitas Timor Timur (Untim) dan sejak 2014 berubah status dari perguruan tinggi swasta menjadi negeri.", "gold"),
             ("📍", "Lokasi Kampus", UNIMOR_ALAMAT, "blue"),
             ("🏅", "Akreditasi", "Terakreditasi Baik oleh BAN-PT (SK No. 1057/SK/BAN-PT/Akred/PT/XII/2020).", "pink")]
    st.markdown('<div class="grid">' + "".join(f'<div class="glass {c}"><div class="ic">{i}</div><h4>{t}</h4><p>{d}</p></div>' for i, t, d, c in cards) + "</div>", unsafe_allow_html=True)
    st.markdown('<div class="sec-title">🎓 Fakultas di UNIMOR</div>', unsafe_allow_html=True)
    cl = ["blue", "pink", "gold", "purple"]
    st.markdown('<div class="grid">' + "".join(f'<div class="glass {cl[n % 4]}"><div class="ic">{i}</div><h4>{t}</h4><p>{d}</p></div>' for n, (i, t, d) in enumerate(UNIMOR_FAKULTAS)) + "</div>", unsafe_allow_html=True)
    st.markdown('<div class="sec-title">📊 UTS • Analisis Data Transaksi Toko</div>', unsafe_allow_html=True)
    flow = '<div class="flow">' + "".join(f'<div class="step"><div class="n">{i}</div><div class="t">{t}</div></div>' for i, t in enumerate(STEPS, 1)) + "</div>"
    st.markdown('<div class="glass gold">'
                f"<h4>{matkul} • EDA, Preprocessing, Apriori, K-Means &amp; Dashboard BI</h4>"
                f"<p><b>Nama:</b> {nama} &nbsp;|&nbsp; <b>NPM:</b> {npm} &nbsp;|&nbsp; <b>{kelas}</b> &nbsp;|&nbsp; <b>{semester}</b></p>"
                f'<div style="margin-top:12px">{flow}</div>'
                "<p>Sistem ini mengubah data transaksi toko menjadi informasi: kualitas data, pola kombinasi produk (Apriori), "
                "segmentasi pelanggan (RFM + K-Means), dashboard interaktif, dan rekomendasi keputusan bisnis. "
                "Klik tombol di bawah untuk mulai.</p></div>", unsafe_allow_html=True)
    st.markdown('<div class="sec-title">🧭 Cara Menggunakan</div>', unsafe_allow_html=True)
    how = [("1️⃣", "Siapkan data", "Letakkan file Excel di folder program atau upload di halaman Analisis.", "blue"),
           ("2️⃣", "Jelajahi 7 tab", "Ikuti urutan tab dari Business Understanding sampai Rekomendasi.", "gold"),
           ("3️⃣", "Atur & filter", "Ubah parameter Apriori, K-Means, dan filter dashboard sesuai kebutuhan.", "pink"),
           ("4️⃣", "Unduh hasil", "Simpan dataset bersih dan seluruh hasil analisis dalam format Excel.", "purple")]
    st.markdown('<div class="grid">' + "".join(f'<div class="glass {c}"><div class="ic">{i}</div><h4>{t}</h4><p>{d}</p></div>' for i, t, d, c in how) + "</div>", unsafe_allow_html=True)
    st.write("")
    _, mid, _ = st.columns([1, 1.2, 1])
    with mid:
        if stretch(st.button, "🚀 Mulai Analisis", type="primary", key="cta"):
            st.session_state.page = "analisis"
            do_rerun()
    st.markdown(f'<div class="foot">© Universitas Timor • {matkul} • {nama} ({npm})</div>', unsafe_allow_html=True)


if st.session_state.page == "beranda":
    render_beranda()
    st.stop()

st.markdown(bg_css(True), unsafe_allow_html=True)
topbar(False)


# ============================================================================
# FUNGSI ANALISIS
# ============================================================================
@st.cache_data(show_spinner="Membaca file Excel...")
def load_excel(file_bytes: bytes):
    xl = pd.ExcelFile(io.BytesIO(file_bytes))
    low = {s.strip().lower(): s for s in xl.sheet_names}
    for need in ["data_raw", "data_produk", "data_pelanggan"]:
        if need not in low:
            raise ValueError(f"Sheet '{need}' tidak ditemukan. Sheet tersedia: {xl.sheet_names}")
    raw = xl.parse(low["data_raw"])
    prod = xl.parse(low["data_produk"])
    pel = xl.parse(low["data_pelanggan"])
    for d in (raw, prod, pel):
        d.columns = [str(c).strip() for c in d.columns]
    return raw, prod, pel


RAW_COLS = ["Transaction_ID", "Transaction_Date", "Customer_ID", "Region", "Customer_Type", "Product_ID",
            "Product_Name", "Category", "Quantity", "Unit_Price", "Discount_Rate", "Line_Total", "Payment_Method"]


def quality_report(d: pd.DataFrame, prod: pd.DataFrame, pel: pd.DataFrame) -> dict:
    """Pemeriksaan kualitas data (dipakai sebelum & sesudah preprocessing)."""
    ref_price = d["Product_ID"].map(prod.set_index("Product_ID")["Reference_Unit_Price_Rp"])
    ref_cat = d["Product_Name"].map(prod.set_index("Product_Name")["Category"])
    ref_pid = d["Product_Name"].map(prod.set_index("Product_Name")["Product_ID"])
    q, p, t, dr = d["Quantity"], d["Unit_Price"], d["Line_Total"], d["Discount_Rate"]
    hit = (q * p * (1 - dr) - t).abs() > 1
    miss = int(d[RAW_COLS].isna().sum().sum() + (d["Payment_Method"] == UNKNOWN).sum())
    return {
        "Jumlah baris": len(d), "Jumlah kolom": len(RAW_COLS),
        "Missing value (sel kosong)": miss,
        "Duplikat persis (13 kolom)": int(d[RAW_COLS].duplicated().sum()),
        "Duplikat kunci (Transaction_ID + Produk)": int(d.duplicated(["Transaction_ID", "Product_Name"]).sum()),
        "Quantity kosong / ≤ 0": int((q.isna() | (q <= 0)).sum()),
        f"Quantity tidak wajar (> {MAX_QTY_WAJAR})": int((q > MAX_QTY_WAJAR).sum()),
        "Unit_Price kosong / ≤ 0": int((p.isna() | (p <= 0)).sum()),
        "Unit_Price tidak sesuai harga referensi": int((p.notna() & (p > 0) & (p != ref_price)).sum()),
        "Line_Total kosong": int(t.isna().sum()),
        "Line_Total ≠ Qty × Harga × (1 − Diskon)": int((hit & t.notna()).sum()),
        "Kategori tidak baku (huruf/spasi)": int((d["Category"].astype("string") != ref_cat.astype("string")).sum()),
        "Metode bayar tidak baku": int((d["Payment_Method"].notna() & (d["Payment_Method"] != UNKNOWN) & ~d["Payment_Method"].isin(PAY_CANON.values())).sum()),
        "Product_ID kosong / tidak cocok nama": int((d["Product_ID"].isna() | (d["Product_ID"] != ref_pid)).sum()),
        "Customer_ID kosong / tidak terdaftar": int((d["Customer_ID"].isna() | ~d["Customer_ID"].isin(pel["Customer_ID"])).sum()),
    }


@st.cache_data(show_spinner="Melakukan data cleaning...")
def clean_data(raw: pd.DataFrame, prod: pd.DataFrame, pel: pd.DataFrame):
    log = []  # (aturan, alasan, temuan, tindakan, dihapus, diperbaiki)
    d = raw[RAW_COLS].copy()
    n0 = len(d)

    # 1. duplikat persis
    n = int(d.duplicated().sum())
    d = d.drop_duplicates().reset_index(drop=True)
    log.append(("Duplikat persis", "Baris identik pada 13 kolom adalah salinan data, bukan pembelian baru. Jika dibiarkan, Frequency, Monetary dan support Apriori membengkak.",
                n, "Hapus, simpan 1 baris", n, 0))

    # 2. Product_ID dari nama produk (master)
    pid = prod.set_index("Product_Name")["Product_ID"]
    bad = d["Product_ID"].isna() | (d["Product_ID"] != d["Product_Name"].map(pid))
    n = int(bad.sum())
    d["Product_ID"] = d["Product_Name"].map(pid)
    log.append(("Product_ID kosong", "Nama produk terisi lengkap dan unik terhadap Data_Produk, sehingga ID dapat dipulihkan tanpa menebak.",
                n, "Isi dari Data_Produk berdasarkan Product_Name", 0, n))

    # 3. Kategori baku
    cat = d["Product_Name"].map(prod.set_index("Product_Name")["Category"])
    n = int((d["Category"] != cat).sum())
    d["Category"] = cat
    log.append(("Kategori tidak baku (sembako / SEMBAKO)", "Perbedaan huruf besar-kecil memecah satu kategori menjadi beberapa label dan merusak agregasi.",
                n, "Samakan dengan kategori resmi Data_Produk", 0, n))

    # 4. Metode pembayaran
    pm = d["Payment_Method"].astype("string").str.strip().str.lower().map(PAY_CANON)
    n_inc = int((d["Payment_Method"].notna() & (pm != d["Payment_Method"])).sum())
    n_na = int(d["Payment_Method"].isna().sum())
    d["Payment_Method"] = pm.astype(object).where(pm.notna(), UNKNOWN)
    log.append(("Metode bayar tidak baku (Qris / qris)", "Variasi ejaan dianggap kategori berbeda. Dibakukan menjadi Tunai, Debit, Transfer, QRIS.",
                n_inc, "Standarisasi teks", 0, n_inc))
    log.append(("Metode bayar kosong", "Metode bayar berbeda antar baris dalam satu nota sehingga tidak dapat diturunkan; tidak dipakai di RFM/Apriori, maka baris dipertahankan.",
                n_na, f"Beri label '{UNKNOWN}'", 0, n_na))

    # 5. Customer_ID kosong -> dari nota yang sama
    miss = d["Customer_ID"].isna()
    n = int(miss.sum())
    fill = d.groupby("Transaction_ID")["Customer_ID"].transform("first")
    d["Customer_ID"] = d["Customer_ID"].fillna(fill)
    d["Customer_ID"] = d["Customer_ID"].astype("string").str.strip().str.upper()
    sisa = d["Customer_ID"].isna() | ~d["Customer_ID"].isin(pel["Customer_ID"])
    n_sisa = int(sisa.sum())
    d = d[~sisa].copy()
    log.append(("Customer_ID kosong / tidak terdaftar", "Satu Transaction_ID hanya milik satu pelanggan, sehingga ID kosong dapat diisi dari baris lain pada nota yang sama. Baris yang tetap tidak dapat diidentifikasi dihapus karena RFM butuh identitas.",
                n, "Isi dari nota yang sama; hapus bila tetap kosong", n_sisa, n - n_sisa))

    # 6. Region & Customer_Type dari master pelanggan
    pm_ = pel.set_index("Customer_ID")
    r_bad = d["Region"].isna() | (d["Region"] != d["Customer_ID"].map(pm_["Region"]))
    t_bad = d["Customer_Type"].isna() | (d["Customer_Type"] != d["Customer_ID"].map(pm_["Customer_Type"]))
    d["Region"] = d["Customer_ID"].map(pm_["Region"])
    d["Customer_Type"] = d["Customer_ID"].map(pm_["Customer_Type"])
    log.append(("Region / Customer_Type kosong", "Atribut pelanggan bersifat tetap sehingga diambil dari master Data_Pelanggan (tidak ada konflik pada data yang terisi).",
                int(r_bad.sum() + t_bad.sum()), "Isi dari Data_Pelanggan", 0, int(r_bad.sum() + t_bad.sum())))

    # 7. Harga satuan
    refp = d["Product_ID"].map(prod.set_index("Product_ID")["Reference_Unit_Price_Rp"]).astype(float)
    pbad = d["Unit_Price"].isna() | (d["Unit_Price"] != refp)
    n = int(pbad.sum())
    d["Unit_Price"] = refp
    log.append(("Unit_Price kosong / negatif / 999.999", "Seluruh harga valid pada data = harga referensi Data_Produk, sehingga nilai -100, 999999 dan kosong jelas salah input.",
                n, "Ganti dengan harga referensi produk", 0, n))

    # 8. Quantity: rekonstruksi dari Line_Total
    q = d["Quantity"]
    qbad = q.isna() | (q <= 0) | (q > MAX_QTY_WAJAR)
    q_rec = (d["Line_Total"] / (d["Unit_Price"] * (1 - d["Discount_Rate"]))).round()
    can = qbad & q_rec.between(1, 6)
    n_bad = int(qbad.sum())
    d["Quantity"] = q.where(~qbad, q_rec)
    unrec = qbad & ~can
    n_un = int(unrec.sum())
    d = d[~unrec].copy()
    d["Quantity"] = d["Quantity"].round().astype(int)
    log.append(("Quantity kosong / ≤ 0 / 999 (outlier)", "Kuantitas tidak mungkin 0 atau negatif, dan 999 adalah kode placeholder. Line_Total pada baris tsb masih benar, sehingga Quantity = Line_Total ÷ (Harga × (1−Diskon)), diterima bila bulat 1–6.",
                n_bad, "Rekonstruksi dari Line_Total; hapus bila tak dapat dipulihkan", n_un, n_bad - n_un))

    # 9. Line_Total
    hitung = d["Quantity"] * d["Unit_Price"] * (1 - d["Discount_Rate"])
    tbad = d["Line_Total"].isna() | ((hitung - d["Line_Total"]).abs() > 1)
    n = int(tbad.sum())
    d["Line_Total"] = hitung.round(0)
    log.append(("Line_Total kosong / tidak konsisten", "Total harus = Quantity × Unit_Price × (1 − Discount_Rate). Dihitung ulang agar Monetary akurat.",
                n, "Hitung ulang dengan rumus", 0, n))

    # 10. Duplikat kunci (nota + produk) setelah perbaikan
    d["_prio"] = d["Payment_Method"].eq(UNKNOWN).astype(int)
    d = d.sort_values(["Transaction_ID", "Product_ID", "_prio"], kind="stable")
    n1 = len(d)
    d = d.drop_duplicates(["Transaction_ID", "Product_ID"], keep="first").drop(columns="_prio")
    n = n1 - len(d)
    log.append(("Duplikat kunci (nota + produk sama, isi berbeda)", "Satu produk hanya muncul sekali per nota. Salinan yang berbeda hanya karena kolom rusak (mis. Quantity 999 atau metode bayar kosong) menjadi identik setelah perbaikan.",
                n, "Simpan 1 baris paling lengkap", n, 0))

    # 11. Transformasi tipe & fitur turunan
    d["Transaction_Date"] = pd.to_datetime(d["Transaction_Date"])
    d["Tanggal"] = d["Transaction_Date"].dt.normalize()
    d["Bulan"] = d["Transaction_Date"].dt.to_period("M").astype(str)
    d["Bulan_No"] = d["Transaction_Date"].dt.month
    d["Jam"] = d["Transaction_Date"].dt.hour
    d["Hari"] = d["Transaction_Date"].dt.day_name()
    d["Diskon_Rp"] = (d["Quantity"] * d["Unit_Price"] - d["Line_Total"]).round(0)
    for c in ["Transaction_ID", "Customer_ID", "Product_ID", "Product_Name", "Category", "Region", "Customer_Type", "Payment_Method"]:
        d[c] = d[c].astype(str)
    d = d.sort_values(["Transaction_Date", "Transaction_ID", "Product_ID"]).reset_index(drop=True)
    log.append(("Transformasi tipe data", "Tanggal → datetime; Quantity → integer; Unit_Price/Line_Total → angka rupiah; ditambah fitur Bulan, Jam, Hari, Diskon_Rp.",
                len(d), "Konversi tipe + fitur turunan", 0, len(d)))

    ld = pd.DataFrame(log, columns=["Aturan", "Alasan pemilihan metode", "Temuan (baris)", "Tindakan", "Baris dihapus", "Baris diperbaiki"])
    ld.insert(0, "No", range(1, len(ld) + 1))
    return d, ld


def build_rfm(d: pd.DataFrame, ref_date: pd.Timestamp, pel: pd.DataFrame) -> pd.DataFrame:
    g = d.groupby("Customer_ID").agg(Last_Transaction=("Tanggal", "max"), Frequency=("Transaction_ID", "nunique"),
                                     Monetary=("Line_Total", "sum")).reset_index()
    g["Recency"] = (ref_date - g["Last_Transaction"]).dt.days
    g = g.merge(pel, on="Customer_ID", how="left")
    g["Log_Frequency"] = np.log(g["Frequency"] + 1)
    g["Log_Monetary"] = np.log(g["Monetary"] + 1)
    for s, t in [("Recency", "Z_Recency"), ("Log_Frequency", "Z_Log_Frequency"), ("Log_Monetary", "Z_Log_Monetary")]:
        g[t] = (g[s] - g[s].mean()) / g[s].std(ddof=0)
    return g


def build_tx(d: pd.DataFrame) -> pd.DataFrame:
    t = d.groupby("Transaction_ID").agg(Customer_ID=("Customer_ID", "first"), Tanggal=("Tanggal", "first"),
                                        Bulan=("Bulan", "first"), basket_size=("Product_ID", "nunique"),
                                        Total_Item=("Quantity", "sum"), Total_Belanja=("Line_Total", "sum")).reset_index()
    return t


def basket_matrix(d: pd.DataFrame):
    m = pd.crosstab(d["Transaction_ID"], d["Product_Name"]).clip(upper=1).astype(bool)
    return m


# ---------------- Apriori (murni numpy) ----------------
@st.cache_data(show_spinner="Menjalankan Apriori...")
def apriori_run(B: np.ndarray, items: tuple, min_sup: float, min_conf: float, max_len: int = 3):
    n = B.shape[0]
    items = list(items)
    sup1 = B.mean(0)
    freq = {(i,): float(sup1[i]) for i in range(len(items)) if sup1[i] >= min_sup}
    level = sorted(freq)
    allf = dict(freq)
    k = 2
    while level and k <= max_len:
        cand = set()
        for a, b in itertools.combinations(level, 2):
            if a[:-1] == b[:-1]:
                c = tuple(sorted(set(a) | set(b)))
                if all(tuple(x) in allf for x in itertools.combinations(c, k - 1)):
                    cand.add(c)
        new = {}
        for c in sorted(cand):
            s = float(B[:, list(c)].all(1).mean())
            if s >= min_sup:
                new[c] = s
        allf.update(new)
        level = sorted(new)
        k += 1
    fi = pd.DataFrame([{"Itemset": " + ".join(items[i] for i in k_), "Ukuran": len(k_), "Support": v,
                        "Frekuensi (transaksi)": int(round(v * n))} for k_, v in allf.items()])
    if fi.empty:
        return fi, pd.DataFrame()
    fi = fi.sort_values(["Ukuran", "Support"], ascending=[True, False]).reset_index(drop=True)
    rows = []
    for iset, s in allf.items():
        if len(iset) < 2:
            continue
        for r in range(1, len(iset)):
            for ant in itertools.combinations(iset, r):
                cons = tuple(x for x in iset if x not in ant)
                conf = s / allf[ant]
                if conf >= min_conf:
                    sc = allf.get(cons, float(B[:, list(cons)].all(1).mean()))
                    rows.append({"Antecedent": " + ".join(items[i] for i in ant), "Consequent": " + ".join(items[i] for i in cons),
                                 "Support": s, "Confidence": conf, "Lift": conf / sc,
                                 "Leverage": s - allf[ant] * sc, "Ukuran": len(iset)})
    rules = pd.DataFrame(rows)
    if not rules.empty:
        rules = rules.sort_values(["Lift", "Confidence", "Support"], ascending=False).reset_index(drop=True)
        rules.insert(0, "Aturan", rules["Antecedent"] + "  →  " + rules["Consequent"])
    return fi, rules


def top_rules(rules: pd.DataFrame, n=10):
    if rules.empty:
        return rules
    r = rules[rules["Lift"] > 1].copy()
    return r.head(n).reset_index(drop=True)


# ---------------- K-Means (murni numpy) ----------------
def _d2(Z, C_):
    return (Z ** 2).sum(1)[:, None] - 2 * Z @ C_.T + (C_ ** 2).sum(1)[None]


def kmeans_np(Z, k, n_init=10, seed=42, max_iter=300):
    rng = np.random.default_rng(seed)
    best = None
    for init_ke in range(1, n_init + 1):
        cent = [Z[rng.integers(len(Z))]]
        for _ in range(1, k):
            d2 = np.maximum(_d2(Z, np.array(cent)).min(1), 0)
            cent.append(Z[rng.choice(len(Z), p=d2 / d2.sum())])
        Cc = np.array(cent, float)
        dist0 = _d2(Z, Cc)
        lab0 = dist0.argmin(1)
        hist = [{"Iterasi": 0, "Inertia": float(np.maximum(dist0.min(1), 0).sum()), "Pindah": np.nan, "Shift": np.nan}]
        labels, conv = None, False
        for it in range(1, max_iter + 1):
            new = _d2(Z, Cc).argmin(1)
            pindah = int((new != (labels if labels is not None else lab0)).sum())
            if labels is not None and pindah == 0:
                conv = True
                hist.append({"Iterasi": it, "Inertia": float(((Z - Cc[new]) ** 2).sum()), "Pindah": 0, "Shift": 0.0})
                labels = new
                break
            labels = new
            old = Cc.copy()
            for j in range(k):
                m = labels == j
                Cc[j] = Z[m].mean(0) if m.any() else Z[rng.integers(len(Z))]
            hist.append({"Iterasi": it, "Inertia": float(((Z - Cc[labels]) ** 2).sum()), "Pindah": pindah,
                         "Shift": float(np.linalg.norm(Cc - old, axis=1).sum())})
        inertia = float(((Z - Cc[labels]) ** 2).sum())
        if best is None or inertia < best["inertia"]:
            best = {"inertia": inertia, "labels": labels.copy(), "centers": Cc.copy(), "iter": it, "hist": hist,
                    "init_ke": init_ke, "converged": conv}
    return best


def silhouette_np(Z, labels):
    D = np.sqrt(np.maximum(_d2(Z, Z), 0))
    ks = np.unique(labels)
    s = np.zeros(len(Z))
    own = np.zeros(len(Z))
    other = np.full(len(Z), np.inf)
    for c in ks:
        m = labels == c
        mean_c = D[:, m].sum(1)
        own[m] = mean_c[m] / max(m.sum() - 1, 1)
        for_other = mean_c / m.sum()
        mask = ~m
        other[mask] = np.minimum(other[mask], for_other[mask])
    s = (other - own) / np.maximum(own, other)
    s[np.array([np.sum(labels == l) for l in labels]) <= 1] = 0
    return float(np.nanmean(s))


@st.cache_data(show_spinner="Menghitung Elbow & Silhouette (k = 2 … 8)...")
def elbow_table(Z, n_init, seed):
    rows = []
    for k in range(2, 9):
        r = kmeans_np(Z, k, n_init, seed)
        rows.append({"k": k, "Inertia": r["inertia"], "Silhouette": silhouette_np(Z, r["labels"]), "Iterasi": r["iter"]})
    t = pd.DataFrame(rows)
    t["Penurunan Inertia"] = -t["Inertia"].diff()
    t["Penurunan (%)"] = t["Penurunan Inertia"] / t["Inertia"].shift(1) * 100
    return t


def label_clusters(centers):
    """centers kolom: [Z_R, Z_logF, Z_logM]. Label berdasar bukti rata-rata Z-score tiap klaster."""
    k = len(centers)
    skor = centers[:, 1] + centers[:, 2] - centers[:, 0]
    out, used = {}, set()
    loyal = int(np.argmax(skor))
    out[loyal] = "Loyal"
    used.add(loyal)
    rest = [j for j in range(k) if j not in used]
    if rest:
        churn = max(rest, key=lambda j: centers[j, 0] - 0.5 * (centers[j, 1] + centers[j, 2]))
        if centers[churn, 0] > 0:
            out[churn] = "Churn"
            used.add(churn)
    for j in range(k):
        if j in used:
            continue
        # Recency rendah + F/M di bawah rata-rata = pelanggan baru; selain itu = potensial
        out[j] = "New" if (centers[j, 0] < 0 and centers[j, 1] + centers[j, 2] < 0) else "Potential"
    return out


LABEL_FULL = {"Loyal": "Loyal Customers", "Potential": "Potential Customers", "New": "New Customers", "Churn": "Churn Risk"}

STRATEGI = {
    "Loyal": ("Pertahankan & tingkatkan nilai", ["Program poin/member tier & akses awal promo.", "Bundling premium berdasar aturan Apriori (cross-sell).", "Program referral ajak pelanggan baru.", "Hindari diskon besar tanpa syarat — mereka sudah rutin membeli."]),
    "Potential": ("Dorong naik kelas ke Loyal", ["Voucher kunjungan berikutnya dengan minimum belanja.", "Rekomendasi produk pelengkap (aturan Apriori).", "Pengingat belanja rutin via WhatsApp.", "Target: naikkan Frequency 1–2 transaksi."]),
    "New": ("Aktivasi & bentuk kebiasaan", ["Welcome voucher untuk pembelian kedua.", "Edukasi produk & katalog lewat WhatsApp.", "Tawarkan paket hemat produk kebutuhan harian.", "Ajak jadi Member agar terikat program loyalitas."]),
    "Churn": ("Win-back selektif & terukur", ["Kampanye \"kami rindu Anda\" + voucher berbatas waktu.", "Survei singkat alasan berhenti belanja.", "Prioritaskan yang Monetary historisnya tinggi.", "Hentikan kontak bila tidak merespons agar biaya iklan efisien."]),
}


def base_label(s):
    for k, v in LABEL_FULL.items():
        if s.startswith(v):
            return k
    return "Potential"


# ============================================================================
# HEADER & UPLOAD
# ============================================================================
foto_up = None
foto_src = get_foto()
inisial = "".join(w[0] for w in nama.split()[:2]).upper()
lg = get_logo()
st.markdown(
    f'<div class="hero">{"<img class=lg src=%s>" % lg if lg else ""}'
    f'{"<img class=pf src=%s>" % foto_src if foto_src else f"<div class=avatar>{inisial}</div>"}<div>'
    f'<span class="badge">UTS • {matkul.upper()}</span>'
    "<h1>📊 Analisis Data Transaksi Toko<br>EDA • Apriori • K-Means • Dashboard BI</h1>"
    '<p class="sub">Business Understanding → Preprocessing → Apriori → Segmentasi RFM → Dashboard → Rekomendasi Keputusan</p>'
    f'<div class="chips"><span>👤 <b>Nama</b> : {nama}</span><span>🆔 <b>NPM</b> : {npm}</span>'
    f"<span>🏫 <b>Kelas</b> : {kelas}</span><span>📘 <b>Semester</b> : {semester}</span></div></div></div>",
    unsafe_allow_html=True)
st.markdown('<div class="tip" style="margin-top:12px">💡 <b>Panduan cepat:</b> (1) data dibaca otomatis atau upload Excel di bawah → '
            '(2) buka <b>7 tab</b> secara berurutan dari kiri ke kanan → (3) atur filter di <b>Tab 6</b> dan parameter di <b>Tab 4 & 5</b> → '
            '(4) unduh hasil di akhir <b>Tab 7</b>.</div>', unsafe_allow_html=True)
st.write("")

default_path = next((APP_DIR / n for n in DEFAULT_FILE_NAMES if (APP_DIR / n).exists()), None)
up = st.file_uploader("📂 Upload file Excel UTS (.xlsx) — sheet Data_Raw, Data_Produk, Data_Pelanggan", type=["xlsx"])
with st.expander("⚙️ Parameter analisis (opsional)"):
    p1, p2, p3 = st.columns(3)
    ref_date = pd.Timestamp(p1.date_input("Tanggal acuan Recency", pd.Timestamp("2025-12-31").date(), help="Hari terakhir periode data (31 Des 2025)."))
    n_init = p2.slider("Jumlah inisialisasi K-Means (n_init)", 3, 30, 10)
    seed = p3.number_input("Random seed", 0, 9999, 42)

if up is not None:
    file_bytes, file_name = up.getvalue(), up.name
elif default_path is not None:
    file_bytes, file_name = default_path.read_bytes(), default_path.name + " (dibaca otomatis dari folder program)"
else:
    st.markdown('<div class="note"><b>📥 Belum ada data.</b> Upload file <b>P8__UTS_Business_Intelligence.xlsx</b> pada kotak di atas '
                "(atau letakkan file itu di folder yang sama dengan program ini). Seluruh analisis berjalan otomatis setelah data tersedia.</div>", unsafe_allow_html=True)
    st.stop()

try:
    raw, prod, pel = load_excel(file_bytes)
    miss_cols = [c for c in RAW_COLS if c not in raw.columns]
    if miss_cols:
        raise ValueError(f"Kolom Data_Raw kurang: {miss_cols}")
except Exception as e:
    st.error(f"Gagal membaca file: {e}")
    st.stop()
st.success(f"✅ Data dimuat: {file_name}")

# ============================================================================
# PIPELINE
# ============================================================================
raw["Transaction_Date"] = pd.to_datetime(raw["Transaction_Date"], errors="coerce")
raw["Customer_ID"] = raw["Customer_ID"].astype("string").str.strip().str.upper().astype(object).where(raw["Customer_ID"].notna(), None)
for _c in ["Region", "Customer_Type", "Product_ID", "Product_Name", "Category", "Payment_Method", "Transaction_ID"]:
    raw[_c] = raw[_c].astype("string").str.strip().astype(object).where(raw[_c].notna(), np.nan)

df, clean_log = clean_data(raw, prod, pel)
q_before = quality_report(raw, prod, pel)
q_after = quality_report(df.assign(Transaction_Date=df["Transaction_Date"]), prod, pel)
tx = build_tx(df)
rfm = build_rfm(df, ref_date, pel)
FEATS = ["Z_Recency", "Z_Log_Frequency", "Z_Log_Monetary"]
Z = rfm[FEATS].values
elbow = elbow_table(Z, int(n_init), int(seed))
best_sil_k = int(elbow.loc[elbow["Silhouette"].idxmax(), "k"])
sil3 = elbow[elbow["k"].between(3, 6)]
k_auto = int(sil3.loc[sil3["Silhouette"].idxmax(), "k"])
k_mode = st.session_state.get("k_mode", "Otomatis")
k_opt = k_auto if k_mode == "Otomatis" else int(st.session_state.get("k_pick", k_auto))

km = kmeans_np(Z, k_opt, int(n_init), int(seed))
lab_map = label_clusters(km["centers"])
skor_c = km["centers"][:, 1] + km["centers"][:, 2] - km["centers"][:, 0]
ordr = ["Loyal", "Potential", "New", "Churn"]
urut = sorted(range(k_opt), key=lambda j: (ordr.index(lab_map[j]), -skor_c[j]))
renum = {o: n for n, o in enumerate(urut)}
rfm["Cluster"] = [renum[x] for x in km["labels"]]
centers = np.array([km["centers"][o] for o in urut])
cnt = {}
cl_name = {}
for newj, o in enumerate(urut):
    b = lab_map[o]
    cnt[b] = cnt.get(b, 0) + 1
    cl_name[newj] = LABEL_FULL[b] + (f" {cnt[b]}" if sum(1 for oo in urut if lab_map[oo] == b) > 1 else "")
rfm["Segmen"] = rfm["Cluster"].map(cl_name)
SEG_LIST = [cl_name[j] for j in range(k_opt)]
SEG_COLORS = {s: SEG_PALETTE[i % len(SEG_PALETTE)] for i, s in enumerate(SEG_LIST)}
sil_opt = silhouette_np(Z, km["labels"])
seg_of = rfm.set_index("Customer_ID")["Segmen"]
df["Segmen"] = df["Customer_ID"].map(seg_of)
tx["Segmen"] = tx["Customer_ID"].map(seg_of)

glob = rfm[["Recency", "Frequency", "Monetary"]].mean()
prof = rfm.groupby("Cluster").agg(Jumlah=("Customer_ID", "count"), Avg_Recency=("Recency", "mean"), Avg_Frequency=("Frequency", "mean"),
                                  Avg_Monetary=("Monetary", "mean"), Total_Revenue=("Monetary", "sum")).reset_index()
prof["Segmen"] = prof["Cluster"].map(cl_name)
prof["Pct_Pelanggan"] = prof["Jumlah"] / prof["Jumlah"].sum()
prof["Pct_Revenue"] = prof["Total_Revenue"] / prof["Total_Revenue"].sum()
prof["Basis"] = prof["Segmen"].map(base_label)
cz = pd.DataFrame(centers, columns=["Z_Recency", "Z_Log_Frequency", "Z_Log_Monetary"])
cz["Cluster"] = range(k_opt)
prof = prof.merge(cz, on="Cluster")


def narasi(r):
    rec = "rendah (baru bertransaksi)" if r.Avg_Recency < glob["Recency"] else "tinggi (lama tidak bertransaksi)"
    fr = "tinggi" if r.Avg_Frequency > glob["Frequency"] else "rendah"
    mo = "tinggi" if r.Avg_Monetary > glob["Monetary"] else "rendah"
    return (f"Recency {rec} = {dec(r.Avg_Recency, 1)} hari (rata-rata semua {dec(glob['Recency'], 1)}); Frequency {fr} = {dec(r.Avg_Frequency, 1)} transaksi "
            f"({dec(glob['Frequency'], 1)}); Monetary {mo} = {rp_short(r.Avg_Monetary)} ({rp_short(glob['Monetary'])}).")


prof["Karakteristik"] = prof.apply(narasi, axis=1)

# basket & Apriori global (parameter diatur di tab 4, disimpan di session_state)
Bm = basket_matrix(df)
ITEMS = tuple(Bm.columns)
Bnp = Bm.values
min_sup = float(st.session_state.get("min_sup", 0.015))
min_conf = float(st.session_state.get("min_conf", 0.20))
fi, rules = apriori_run(Bnp, ITEMS, min_sup, min_conf, 3)
top10 = top_rules(rules, 10)

total_rev, n_tx, n_cust = float(df["Line_Total"].sum()), int(df["Transaction_ID"].nunique()), int(df["Customer_ID"].nunique())
avg_basket = float(tx["basket_size"].mean())

tabs = st.tabs(["🧭 1. Business Understanding", "🔍 2. EDA", "🧹 3. Preprocessing & Fitur", "🛒 4. Apriori",
                "👥 5. K-Means (RFM)", "📊 6. Dashboard BI", "🎯 7. Rekomendasi Keputusan"])

# ============================== TAB 1 =========================================
with tabs[0]:
    st.markdown("### 🧭 Tahap 1 • Business Understanding (Soal 1)")
    flow = [("Business Understanding", "Masalah & pertanyaan bisnis"), ("EDA", "Kondisi data awal"), ("Preprocessing", "Cleaning & fitur"),
            ("Apriori", "Pola kombinasi produk"), ("K-Means", "Segmentasi RFM"), ("Dashboard", "KPI interaktif"), ("Rekomendasi", "Keputusan terukur")]
    st.markdown('<div class="flow">' + "".join(f'<div class="step"><div class="n">{i}</div><div class="t">{t}</div><div class="d">{d}</div></div>' for i, (t, d) in enumerate(flow, 1)) + "</div>", unsafe_allow_html=True)
    a, b = st.columns(2)
    with a:
        sec("a. Permasalahan utama perusahaan")
        st.markdown(
            "Berdasarkan narasi studi kasus, ada empat masalah yang saling terkait:\n"
            "1. **Data hanya dipakai untuk laporan rutin** — belum diolah menjadi pengetahuan tentang perilaku belanja.\n"
            "2. **Pola pembelian produk belum diketahui** — manajemen belum tahu produk mana yang dibeli bersamaan sehingga penataan rak, bundling, dan promosi silang masih berdasarkan perkiraan.\n"
            "3. **Pelanggan diperlakukan sama** — tanpa segmentasi, promosi yang sama dikirim ke pelanggan loyal maupun yang hampir berhenti belanja.\n"
            "4. **Kualitas data rendah** — nilai kosong, duplikat, kuantitas/harga tidak wajar, dan teks tidak konsisten membuat laporan menyesatkan jika tidak dibersihkan.")
    with b:
        sec("b. Pertanyaan bisnis (dijawab pada sistem ini)")
        qb = pd.DataFrame({"No": [1, 2, 3, 4, 5], "Pertanyaan bisnis": [
            "Bagaimana tren penjualan bulanan, dan produk/kategori apa yang menyumbang pendapatan terbesar?",
            "Kombinasi produk apa yang sering dibeli bersamaan dan seberapa kuat hubungannya (support, confidence, lift)?",
            "Bagaimana pelanggan dapat dikelompokkan berdasarkan Recency, Frequency, Monetary, dan siapa yang berisiko churn?",
            "Segmen mana yang harus diprioritaskan dan strategi pemasaran apa yang tepat untuk tiap segmen?",
            "Bagaimana perbedaan perilaku belanja antara Member dan Non-Member serta antarwilayah?"],
            "Dijawab di": ["Tab 2 & 6", "Tab 4", "Tab 5", "Tab 5 & 7", "Tab 2 & 6"]})
        show_df(qb)
    sec("c. Bagaimana analisis membantu meningkatkan penjualan & mempertahankan pelanggan")
    h1, h2, h3 = st.columns(3)
    h1.markdown('<div class="info"><b>📈 Meningkatkan penjualan</b><br>Aturan Apriori dipakai untuk bundling, cross-selling, dan penataan rak berdekatan → menaikkan rata-rata basket_size dan nilai per transaksi.</div>', unsafe_allow_html=True)
    h2.markdown('<div class="ok"><b>🤝 Mempertahankan pelanggan</b><br>Segmentasi RFM mengidentifikasi pelanggan loyal untuk dijaga dan pelanggan berisiko churn untuk di-win-back sebelum hilang.</div>', unsafe_allow_html=True)
    h3.markdown('<div class="note"><b>💰 Efisiensi anggaran</b><br>Promosi diarahkan ke segmen yang tepat dengan penawaran berbeda, sehingga biaya iklan tidak terbuang untuk kampanye massal.</div>', unsafe_allow_html=True)
    st.write("")
    sec("Ringkasan dataset yang dianalisis")
    c1, c2, c3, c4, c5 = st.columns(5)
    c1.markdown(kpi("Baris data mentah", angka(len(raw)), "baris item transaksi", "blue"), unsafe_allow_html=True)
    c2.markdown(kpi("Transaksi (nota)", angka(raw["Transaction_ID"].nunique()), "Transaction_ID unik", "pink"), unsafe_allow_html=True)
    c3.markdown(kpi("Pelanggan terdaftar", angka(len(pel)), "Data_Pelanggan", "green"), unsafe_allow_html=True)
    c4.markdown(kpi("Produk", angka(len(prod)), "Data_Produk", "orange"), unsafe_allow_html=True)
    c5.markdown(kpi("Periode", f"{raw['Transaction_Date'].min():%d %b %Y}", f"s.d. {raw['Transaction_Date'].max():%d %b %Y}", "purple"), unsafe_allow_html=True)
    st.markdown('<div class="warn" style="margin-top:12px"><b>Peringatan:</b> satu Transaction_ID memuat beberapa baris produk, jadi jumlah transaksi dihitung dari <b>ID unik</b>, bukan jumlah baris.</div>', unsafe_allow_html=True)

# ============================== TAB 2 =========================================
with tabs[1]:
    st.markdown("### 🔍 Tahap 1 • Exploratory Data Analysis — data MENTAH (Soal 2)")
    sec("a. Struktur data: baris, kolom, tipe data, rentang waktu")
    c1, c2, c3, c4 = st.columns(4)
    c1.markdown(kpi("Jumlah baris", angka(len(raw)), "data mentah", "blue"), unsafe_allow_html=True)
    c2.markdown(kpi("Jumlah kolom", str(raw.shape[1]), "atribut", "pink"), unsafe_allow_html=True)
    c3.markdown(kpi("Tanggal awal", f"{raw['Transaction_Date'].min():%d-%m-%Y}", f"{raw['Transaction_Date'].min():%H:%M}", "green"), unsafe_allow_html=True)
    c4.markdown(kpi("Tanggal akhir", f"{raw['Transaction_Date'].max():%d-%m-%Y}", f"{(raw['Transaction_Date'].max() - raw['Transaction_Date'].min()).days + 1} hari", "orange"), unsafe_allow_html=True)
    ket = {"Transaction_ID": "ID nota", "Transaction_Date": "Waktu transaksi", "Customer_ID": "ID pelanggan", "Region": "Wilayah", "Customer_Type": "Member / Non-Member",
           "Product_ID": "Kode produk", "Product_Name": "Nama produk", "Category": "Kategori", "Quantity": "Jumlah barang", "Unit_Price": "Harga satuan",
           "Discount_Rate": "Diskon", "Line_Total": "Total baris", "Payment_Method": "Metode bayar"}
    info = pd.DataFrame({"Kolom": RAW_COLS, "Tipe data": [str(raw[c].dtype) for c in RAW_COLS], "Keterangan": [ket[c] for c in RAW_COLS],
                         "Nilai unik": [int(raw[c].nunique()) for c in RAW_COLS], "Missing": [int(raw[c].isna().sum()) for c in RAW_COLS],
                         "Missing (%)": [raw[c].isna().mean() * 100 for c in RAW_COLS]})
    st.write("")
    show_df(info, column_config={"Missing (%)": st.column_config.NumberColumn(format="%.2f")})

    sec("b. Statistik deskriptif")
    valid = raw[(raw["Quantity"].between(1, MAX_QTY_WAJAR)) & (raw["Unit_Price"] > 0) & (raw["Unit_Price"] < 100000)]
    d1, d2, d3, d4, d5, d6 = st.columns(6)
    d1.markdown(kpi("Transaksi unik", angka(raw["Transaction_ID"].nunique()), "Transaction_ID", "blue"), unsafe_allow_html=True)
    d2.markdown(kpi("Pelanggan unik", angka(raw["Customer_ID"].nunique()), "pada data mentah", "pink"), unsafe_allow_html=True)
    d3.markdown(kpi("Produk unik", angka(raw["Product_Name"].nunique()), "Product_Name", "green"), unsafe_allow_html=True)
    d4.markdown(kpi("Rata-rata Quantity", dec(valid["Quantity"].mean(), 2), "baris valid saja", "orange"), unsafe_allow_html=True)
    d5.markdown(kpi("Total penjualan", rp_short(raw["Line_Total"].sum()), "Σ Line_Total mentah", "purple"), unsafe_allow_html=True)
    d6.markdown(kpi("Item / nota", dec(raw.drop_duplicates()["Transaction_ID"].value_counts().mean(), 2), "baris per nota", "gold"), unsafe_allow_html=True)
    st.write("")
    desc = raw[["Quantity", "Unit_Price", "Discount_Rate", "Line_Total"]].describe().T.reset_index().rename(columns={"index": "Variabel"})
    show_df(desc, column_config={c: st.column_config.NumberColumn(format="%.2f") for c in ["mean", "std", "min", "25%", "50%", "75%", "max"]})
    st.caption("Perhatikan Quantity min −2 / max 999 dan Unit_Price min −100 / max 999.999: nilai tidak wajar yang menaikkan std secara semu.")

    sec("c. Missing value, duplikat, dan outlier")
    m1, m2 = st.columns([1, 1.2])
    with m1:
        mv = info[info["Missing"] > 0][["Kolom", "Missing", "Missing (%)"]].sort_values("Missing", ascending=False)
        fig = px.bar(mv, x="Missing", y="Kolom", orientation="h", text="Missing", title="Missing value per kolom", color_discrete_sequence=[C["pink"]])
        fig.update_yaxes(autorange="reversed")
        show_chart(fig, 360)
    with m2:
        iqr_rows = []
        for c in ["Quantity", "Unit_Price", "Line_Total"]:
            s = raw[c].dropna()
            q1, q3 = s.quantile(.25), s.quantile(.75)
            lo, hi = q1 - 1.5 * (q3 - q1), q3 + 1.5 * (q3 - q1)
            iqr_rows.append({"Variabel": c, "Q1": q1, "Q3": q3, "Batas bawah IQR": lo, "Batas atas IQR": hi,
                             "Jumlah outlier (IQR)": int(((s < lo) | (s > hi)).sum())})
        show_df(pd.DataFrame(iqr_rows), column_config={c: st.column_config.NumberColumn(format="%.1f") for c in ["Q1", "Q3", "Batas bawah IQR", "Batas atas IQR"]})
        o = pd.DataFrame({"Jenis masalah": ["Duplikat persis", "Quantity ≤ 0", "Quantity 999 (> 20)", "Unit_Price ≤ 0", "Unit_Price 999.999"],
                          "Baris": [int(raw.duplicated().sum()), int((raw["Quantity"] <= 0).sum()), int((raw["Quantity"] > MAX_QTY_WAJAR).sum()),
                                    int((raw["Unit_Price"] <= 0).sum()), int((raw["Unit_Price"] > 100000).sum())]})
        fig = px.bar(o, x="Jenis masalah", y="Baris", text="Baris", title="Duplikat & nilai tidak wajar", color_discrete_sequence=[C["orange"]])
        show_chart(fig, 250)
    st.markdown('<div class="info"><b>Catatan IQR:</b> outlier IQR pada Line_Total banyak berasal dari produk mahal yang wajar (Beras 5 kg, Susu Formula, Popok). '
                'Jadi hanya nilai <b>mustahil secara bisnis</b> (Quantity ≤ 0 / 999, harga negatif / 999.999) yang diperbaiki; nilai mahal-tapi-wajar dipertahankan dan diredam lewat log transform pada RFM.</div>', unsafe_allow_html=True)

    sec("d. Visualisasi awal")
    v1, v2 = st.columns(2)
    with v1:
        rr = raw.drop_duplicates().dropna(subset=["Transaction_Date"]).copy()
        rr["Bulan"] = rr["Transaction_Date"].dt.to_period("M").astype(str)
        mt = rr.groupby("Bulan").agg(Nota=("Transaction_ID", "nunique")).reset_index()
        fig = px.area(mt, x="Bulan", y="Nota", markers=True, title="Jumlah transaksi (nota) per bulan", color_discrete_sequence=[C["indigo"]])
        show_chart(fig, 340)
    with v2:
        tp = raw.drop_duplicates()["Product_Name"].value_counts().head(10).iloc[::-1].reset_index()
        tp.columns = ["Produk", "Frekuensi"]
        fig = px.bar(tp, x="Frekuensi", y="Produk", orientation="h", title="10 produk paling sering dibeli (frekuensi baris)", text="Frekuensi", color_discrete_sequence=[C["teal"]])
        show_chart(fig, 340)
    v3, v4 = st.columns(2)
    with v3:
        qv = raw["Quantity"].fillna(-99).clip(upper=12).round().astype(int).value_counts().sort_index().reset_index()
        qv.columns = ["Quantity", "Baris"]
        qv["Quantity"] = qv["Quantity"].astype(str).replace({"-99": "Kosong", "12": "≥12 (999)"})
        fig = px.bar(qv, x="Quantity", y="Baris", title="Distribusi Quantity (menampilkan nilai aneh)", text="Baris", color_discrete_sequence=[C["red"]])
        fig.update_yaxes(type="log")
        show_chart(fig, 340)
    with v4:
        cat_raw = raw["Category"].astype("string").str.title().value_counts().reset_index()
        cat_raw.columns = ["Kategori", "Baris"]
        fig = px.pie(cat_raw, names="Kategori", values="Baris", hole=.5, title="Proporsi baris per kategori", color_discrete_sequence=CAT_COLORS)
        show_chart(fig, 340)

    sec("e. Temuan awal")
    dup_n = int(raw.duplicated().sum())
    pay_inc = q_before["Metode bayar tidak baku"]
    top_cat = raw.drop_duplicates()["Category"].astype("string").str.title().value_counts().index[0]
    st.markdown(
        f"1. **Data tidak bersih.** Dari {angka(len(raw))} baris, terdapat **{angka(dup_n)} duplikat persis**, {angka(int(raw.isna().sum().sum()))} sel kosong, serta "
        f"Quantity tidak valid pada {angka(q_before['Quantity kosong / ≤ 0'] + q_before[f'Quantity tidak wajar (> {MAX_QTY_WAJAR})'])} baris dan harga tidak valid pada {angka(q_before['Unit_Price kosong / ≤ 0'])} baris.\n"
        f"2. **Satu nota = banyak baris.** {angka(len(raw))} baris hanya berasal dari **{angka(raw['Transaction_ID'].nunique())} nota** (rata-rata {dec(raw.drop_duplicates()['Transaction_ID'].value_counts().mean(), 2)} produk per nota), "
        "sehingga data cocok untuk Market Basket Analysis, dan transaksi harus dihitung dari ID unik.\n"
        f"3. **Penjualan merata sepanjang 2025** dengan kategori terbanyak pada baris: **{top_cat}**; tidak ada satu produk yang mendominasi, jadi pola kombinasi produk perlu dicari lewat Apriori.\n"
        f"4. **Inkonsistensi teks** pada kategori (sembako/SEMBAKO) dan metode bayar (Qris/qris) = {angka(q_before['Kategori tidak baku (huruf/spasi)'] + pay_inc)} baris yang akan memecah agregasi bila tidak dibakukan.\n"
        "5. **Nilai ekstrem bersifat kesalahan input** (Quantity 999, harga 999.999 / −100), bukan perilaku belanja nyata, sehingga layak diperbaiki, bukan dibiarkan.")

# ============================== TAB 3 =========================================
with tabs[2]:
    st.markdown("### 🧹 Tahap 2 • Data Preprocessing & Feature Engineering (Soal 3 & 4)")
    dihapus = len(raw) - len(df)
    c1, c2, c3, c4 = st.columns(4)
    c1.markdown(kpi("Baris mentah", angka(len(raw)), "sebelum", "red"), unsafe_allow_html=True)
    c2.markdown(kpi("Baris bersih", angka(len(df)), "sesudah", "green"), unsafe_allow_html=True)
    c3.markdown(kpi("Baris dihapus", angka(dihapus), f"{dihapus / len(raw):.1%} dari mentah", "orange"), unsafe_allow_html=True)
    c4.markdown(kpi("Baris diperbaiki", angka(int(clean_log["Baris diperbaiki"].iloc[:-1].sum())), "nilai dipulihkan, bukan dibuang", "blue"), unsafe_allow_html=True)
    st.write("")
    sec("Soal 3a–b • Penanganan masalah & alasan pemilihan metode")
    show_df(clean_log, column_config={"Alasan pemilihan metode": st.column_config.TextColumn(width="large"), "Tindakan": st.column_config.TextColumn(width="medium")})

    sec("Soal 3d • Penanganan nilai ekstrem (konteks bisnis)")
    lt = df["Line_Total"]
    q1, q3 = lt.quantile(.25), lt.quantile(.75)
    hi = q3 + 1.5 * (q3 - q1)
    n_hi = int((lt > hi).sum())
    st.markdown(f"- **Nilai mustahil** (Quantity 0/−2/999, harga −100/999.999) → **diperbaiki** dari Line_Total & harga referensi karena merupakan kesalahan input.\n"
                f"- **Nilai mahal tapi wajar**: {angka(n_hi)} baris ({n_hi / len(df):.1%}) melebihi batas atas IQR Line_Total ({rupiah(hi)}), seluruhnya produk bernilai tinggi dengan Quantity ≤ 6. "
                "**Dipertahankan** karena merupakan pembelian nyata (membuang berarti menghilangkan pelanggan bernilai tinggi); dampaknya pada K-Means diredam dengan transformasi log + standardisasi.")

    sec("Soal 3e • Perbandingan sebelum & sesudah preprocessing")
    cmp_rows = []
    for i, (k_, v) in enumerate(q_before.items(), 1):
        a_ = q_after[k_]
        st_ = "ℹ️ Info" if k_ in ("Jumlah baris", "Jumlah kolom") else ("✅ Bersih" if a_ == 0 else "⚠️ Sisa")
        cmp_rows.append({"No": i, "Pemeriksaan": k_, "Sebelum": v, "Sesudah": a_, "Status": st_})
    cmp_df = pd.DataFrame(cmp_rows)
    ca, cb = st.columns([1.1, 1])
    with ca:
        show_df(cmp_df)
        n_unk = int((df["Payment_Method"] == UNKNOWN).sum())
        st.caption(f"Sisa 'Missing value' sesudah cleaning = {n_unk} baris metode bayar berlabel '{UNKNOWN}' (sengaja dipertahankan, tidak dipakai analisis).")
    with cb:
        mv_ = cmp_df[~cmp_df["Pemeriksaan"].isin(["Jumlah baris", "Jumlah kolom"])].melt(id_vars="Pemeriksaan", value_vars=["Sebelum", "Sesudah"], var_name="Tahap", value_name="Jumlah")
        fig = px.bar(mv_, y="Pemeriksaan", x="Jumlah", color="Tahap", orientation="h", barmode="group", title="Temuan masalah: sebelum vs sesudah",
                     color_discrete_map={"Sebelum": C["red"], "Sesudah": C["green"]})
        fig.update_yaxes(autorange="reversed")
        show_chart(fig, 520)

    sec("Soal 4g • Dokumentasi jumlah data dipertahankan / dihapus / diubah")
    rem = clean_log[clean_log["Baris dihapus"] > 0]
    wf = go.Figure(go.Waterfall(x=["Baris mentah"] + [a.split(" (")[0] for a in rem["Aturan"]] + ["Baris bersih"],
                                measure=["absolute"] + ["relative"] * len(rem) + ["total"],
                                y=[len(raw)] + [-int(v) for v in rem["Baris dihapus"]] + [len(df)],
                                decreasing={"marker": {"color": C["red"]}}, totals={"marker": {"color": C["green"]}},
                                text=[angka(len(raw))] + [f"-{angka(v)}" for v in rem["Baris dihapus"]] + [angka(len(df))], textposition="outside"))
    wf.update_layout(title="Alur jumlah baris: mentah → bersih", showlegend=False)
    wf.update_yaxes(range=[len(df) - 1500, len(raw) + 800])
    d1, d2 = st.columns([1.3, 1])
    with d1:
        show_chart(wf, 400)
    with d2:
        doc = pd.DataFrame({"Kategori": ["Dipertahankan", "Dihapus", "Diubah/diperbaiki (nilai)"],
                            "Baris": [len(df), dihapus, int(clean_log["Baris diperbaiki"].iloc[:-1].sum())]})
        fig = px.pie(doc.iloc[:2], names="Kategori", values="Baris", hole=.55, title="Proporsi baris", color_discrete_sequence=[C["green"], C["red"]])
        show_chart(fig, 400)

    st.markdown("---")
    sec("Soal 4a • Fitur RFM per pelanggan")
    st.markdown(f'<div class="info"><b>Recency</b> = {ref_date:%d %b %Y} − tanggal transaksi terakhir (hari) &nbsp;|&nbsp; <b>Frequency</b> = jumlah transaksi (nota) unik &nbsp;|&nbsp; <b>Monetary</b> = Σ Line_Total (Rp). '
                f"Dihitung dari <b>{angka(len(rfm))}</b> pelanggan aktif.</div>", unsafe_allow_html=True)
    st.write("")
    r1, r2, r3 = st.columns(3)
    for col, (v, ttl, clr) in zip([r1, r2, r3], [("Recency", "Distribusi Recency (hari)", C["indigo"]), ("Frequency", "Distribusi Frequency", C["orange"]), ("Monetary", "Distribusi Monetary (Rp)", C["pink"])]):
        with col:
            fig = px.histogram(rfm, x=v, nbins=30, title=ttl, color_discrete_sequence=[clr])
            fig.update_layout(bargap=.05, yaxis_title="Pelanggan")
            show_chart(fig, 280)
    show_df(rfm[["Customer_ID", "Last_Transaction", "Recency", "Frequency", "Monetary", "Region", "Customer_Type", "Age_Group"]].sort_values("Monetary", ascending=False),
            column_config={"Monetary": st.column_config.NumberColumn(format="%d"), "Last_Transaction": st.column_config.DateColumn(format="YYYY-MM-DD")})

    sec("Soal 4b • Fitur basket_size (jumlah produk berbeda per transaksi)")
    b1, b2 = st.columns([1, 1.4])
    with b1:
        c_ = st.columns(2)
        c_[0].markdown(kpi("Rata-rata basket_size", dec(avg_basket, 2), "produk berbeda / nota", "teal" if False else "green"), unsafe_allow_html=True)
        c_[1].markdown(kpi("Maksimum", str(int(tx["basket_size"].max())), "produk / nota", "orange"), unsafe_allow_html=True)
        show_df(tx.head(200))
    with b2:
        bs = tx["basket_size"].value_counts().sort_index().reset_index()
        bs.columns = ["basket_size", "Transaksi"]
        fig = px.bar(bs, x="basket_size", y="Transaksi", text="Transaksi", title="Distribusi basket_size", color_discrete_sequence=[C["purple"]])
        fig.update_xaxes(dtick=1)
        show_chart(fig, 330)

    sec("Soal 4c • Format transaksi untuk Apriori (matriks transaksi × produk, boolean)")
    st.markdown(f"Matriks berukuran **{angka(Bm.shape[0])} transaksi × {Bm.shape[1]} produk**; nilai `True` = produk ada di nota. Kuantitas diabaikan (hanya ada/tidak) karena Apriori menganalisis **kehadiran** item.")
    show_df(Bm.head(15).astype(int).reset_index())

    sec("Soal 4d–e • Matriks fitur K-Means & standardisasi")
    sk = pd.DataFrame({"Fitur": ["Frequency", "Monetary"], "Skewness sebelum log": [rfm["Frequency"].skew(), rfm["Monetary"].skew()],
                       "Skewness sesudah log": [rfm["Log_Frequency"].skew(), rfm["Log_Monetary"].skew()]})
    st.latex(r"Log\_F=\ln(Frequency+1)\quad Log\_M=\ln(Monetary+1)\quad Z=\frac{x-\mu}{\sigma}")
    e1, e2 = st.columns([1, 1.3])
    with e1:
        show_df(sk, column_config={c: st.column_config.NumberColumn(format="%.3f") for c in sk.columns[1:]})
        ver = pd.DataFrame({"Fitur": FEATS, "Mean": [rfm[c].mean() for c in FEATS], "Std Dev": [rfm[c].std(ddof=0) for c in FEATS]})
        ver["Status"] = np.where((ver["Mean"].abs() < 1e-9) & ((ver["Std Dev"] - 1).abs() < 1e-9), "✅ Valid", "⚠️")
        show_df(ver, column_config={"Mean": st.column_config.NumberColumn(format="%.6f"), "Std Dev": st.column_config.NumberColumn(format="%.6f")})
    with e2:
        fg = make_subplots(rows=2, cols=2, subplot_titles=("Frequency (asli)", "Log_Frequency", "Monetary (asli)", "Log_Monetary"))
        for (r_, c_), (col_, clr) in zip([(1, 1), (1, 2), (2, 1), (2, 2)], [("Frequency", C["orange"]), ("Log_Frequency", C["green"]), ("Monetary", C["pink"]), ("Log_Monetary", C["cyan"])]):
            fg.add_trace(go.Histogram(x=rfm[col_], marker_color=clr, nbinsx=30, showlegend=False), row=r_, col=c_)
        fg.update_layout(title="Efek log transformation", bargap=.05)
        show_chart(fg, 400)
    show_df(rfm[["Customer_ID", "Recency", "Frequency", "Monetary", "Log_Frequency", "Log_Monetary"] + FEATS],
            column_config={c: st.column_config.NumberColumn(format="%.4f") for c in ["Log_Frequency", "Log_Monetary"] + FEATS} | {"Monetary": st.column_config.NumberColumn(format="%d")})
    qa("Mengapa memakai Z-score (standardisasi) dan bukan Min-Max?",
       "K-Means memakai jarak Euclidean, sehingga fitur berskala besar (Monetary jutaan Rupiah) akan mendominasi Recency (hari). **Z-score** menyamakan skala tiap fitur (mean 0, std 1) "
       "dan tidak terlalu terpengaruh oleh satu nilai maksimum seperti Min-Max. Log dilakukan lebih dulu karena Frequency & Monetary menceng ke kanan; "
       f"skewness Monetary turun dari **{sk.iloc[1, 1]:.2f}** menjadi **{sk.iloc[1, 2]:.2f}**.")
    qa("Soal 4f • Strategi mengolah dataset berukuran besar",
       "- **Agregasi:** ubah data level baris menjadi level pelanggan (RFM) / level nota (basket) — dilakukan pada sistem ini sehingga K-Means hanya memproses ±1.500 titik, bukan 18 ribu baris.\n"
       "- **Chunking:** baca file per potongan (`pd.read_csv(chunksize=...)`) lalu agregasi bertahap dan gabungkan hasilnya.\n"
       "- **Sampling:** untuk jutaan transaksi, ambil sampel acak/berstrata (mis. per bulan atau wilayah) untuk eksplorasi dan tuning k.\n"
       "- **Pruning Apriori:** naikkan min_support, batasi panjang itemset (di sini ≤ 3), dan pakai matriks boolean/sparse agar memori hemat; FP-Growth bila itemset sangat banyak.\n"
       "- **Tipe data hemat memori:** `category`, `int32/float32`, serta caching hasil (`st.cache_data`).")

# ============================== TAB 4 =========================================
with tabs[3]:
    st.markdown("### 🛒 Tahap 3 • Market Basket Analysis — Apriori (Soal 5 & 6)")
    st.markdown("#### 5a–b • Matriks transaksi & penentuan parameter")
    pa, pb = st.columns(2)
    with pa:
        st.slider("Minimum support", 0.005, 0.10, 0.015, 0.005, key="min_sup", format="%.3f",
                  help="Persentase minimum nota yang memuat kombinasi produk.")
    with pb:
        st.slider("Minimum confidence", 0.05, 0.80, 0.20, 0.05, key="min_conf", format="%.2f",
                  help="Peluang consequent dibeli jika antecedent dibeli.")
    avg_sup1 = Bm.mean().mean()
    st.markdown(
        f'<div class="info"><b>Alasan pemilihan nilai (default support = 1,5%, confidence = 20%):</b> '
        f"Dataset berisi {angka(Bm.shape[0])} nota dengan {Bm.shape[1]} produk dan rata-rata {dec(avg_basket, 2)} produk/nota. Support satu produk rata-rata hanya {avg_sup1:.1%}, "
        f"sehingga support pasangan secara acak ≈ {avg_sup1 ** 2:.2%}. Support 1,5% (≈ {angka(0.015 * Bm.shape[0])} nota) cukup tinggi untuk menyaring kebetulan, tetapi cukup rendah agar pola nyata tidak hilang "
        "(support terlalu besar → tidak ada aturan; terlalu kecil → ratusan aturan acak). Confidence 20% = minimal 1 dari 5 pembeli antecedent juga membeli consequent, cukup bermakna bagi tindakan bundling. "
        "Filter tambahan: <b>lift > 1</b> agar hanya hubungan yang lebih kuat dari kebetulan yang diambil.</div>", unsafe_allow_html=True)
    st.write("")
    c1, c2, c3, c4 = st.columns(4)
    c1.markdown(kpi("Matriks transaksi", f"{angka(Bm.shape[0])} × {Bm.shape[1]}", "nota × produk", "blue"), unsafe_allow_html=True)
    c2.markdown(kpi("Frequent itemset", angka(len(fi)), f"support ≥ {min_sup:.1%}", "pink"), unsafe_allow_html=True)
    c3.markdown(kpi("Aturan asosiasi", angka(len(rules)), f"confidence ≥ {min_conf:.0%}", "green"), unsafe_allow_html=True)
    c4.markdown(kpi("Aturan lift > 1", angka(int((rules["Lift"] > 1).sum())) if not rules.empty else "0", "hubungan positif", "orange"), unsafe_allow_html=True)
    st.write("")
    with st.expander("Lihat matriks transaksi (contoh 30 nota pertama)"):
        show_df(Bm.head(30).astype(int).reset_index())

    st.markdown("#### 5c • Frequent itemsets")
    f1, f2 = st.columns([1, 1.2])
    with f1:
        show_df(fi, column_config={"Support": st.column_config.NumberColumn(format="%.4f")})
    with f2:
        f1_ = fi[fi["Ukuran"] == 1].sort_values("Support").tail(15)
        fig = px.bar(f1_, x="Support", y="Itemset", orientation="h", title="Support produk tunggal (15 teratas)", color_discrete_sequence=[C["cyan"]])
        show_chart(fig, 430)

    if rules.empty or top10.empty:
        st.warning("Tidak ada aturan yang memenuhi parameter. Turunkan minimum support / confidence.")
    else:
        st.markdown("#### 5d–f • Aturan asosiasi & 10 aturan terbaik")
        st.markdown("**Kriteria peringkat:** lift tertinggi (kekuatan hubungan di atas kebetulan), lalu confidence, lalu support. Hanya aturan dengan **lift > 1**.")
        show_df(top10[["Aturan", "Support", "Confidence", "Lift", "Leverage"]],
                column_config={"Support": st.column_config.NumberColumn(format="%.4f"), "Confidence": st.column_config.NumberColumn(format="%.3f"),
                               "Lift": st.column_config.NumberColumn(format="%.3f"), "Leverage": st.column_config.NumberColumn(format="%.4f")})
        with st.expander(f"Semua aturan ({angka(len(rules))})"):
            show_df(rules.drop(columns=["Ukuran"]), column_config={c: st.column_config.NumberColumn(format="%.4f") for c in ["Support", "Confidence", "Lift", "Leverage"]})

        st.markdown("#### 5g • Visualisasi Apriori")
        g1, g2 = st.columns(2)
        with g1:  # network graph (layout lingkaran, tanpa networkx)
            nodes = sorted(set(top10["Antecedent"]) | set(top10["Consequent"]))
            ang = {n: 2 * np.pi * i / len(nodes) for i, n in enumerate(nodes)}
            pos = {n: (np.cos(a), np.sin(a)) for n, a in ang.items()}
            fig = go.Figure()
            lmin, lmax = top10["Lift"].min(), top10["Lift"].max()
            for _, r in top10.iterrows():
                x0, y0 = pos[r["Antecedent"]]
                x1, y1 = pos[r["Consequent"]]
                w = 1.5 + 7 * (r["Lift"] - lmin) / max(lmax - lmin, 1e-9)
                fig.add_trace(go.Scatter(x=[x0, x1], y=[y0, y1], mode="lines", line=dict(width=w, color=px.colors.sample_colorscale("Turbo", [(r["Lift"] - lmin) / max(lmax - lmin, 1e-9) * .8 + .1])[0]),
                                         hoverinfo="text", text=f"{r['Aturan']}<br>Lift {r['Lift']:.2f} • Conf {r['Confidence']:.0%}", showlegend=False))
                fig.add_annotation(x=x1, y=y1, ax=x0, ay=y0, xref="x", yref="y", axref="x", ayref="y", showarrow=True, arrowhead=3, arrowsize=1.1, arrowwidth=1, arrowcolor="rgba(0,0,0,.35)", standoff=14)
            supmap = dict(zip(ITEMS, Bm.mean()))
            fig.add_trace(go.Scatter(x=[pos[n][0] for n in nodes], y=[pos[n][1] for n in nodes], mode="markers+text", text=nodes, textposition=["top center" if pos[n][1] > 0 else "bottom center" for n in nodes],
                                     marker=dict(size=[18 + 200 * supmap.get(n, .1) for n in nodes], color=C["purple"], line=dict(color="#fff", width=2)), showlegend=False, hoverinfo="text",
                                     hovertext=[f"{n}<br>Support {supmap.get(n, 0):.1%}" for n in nodes]))
            fig.update_layout(title="Network graph aturan terbaik (tebal garis = lift, panah = antecedent → consequent)", xaxis=dict(visible=False, range=[-1.6, 1.6]), yaxis=dict(visible=False, range=[-1.4, 1.4]))
            show_chart(fig, 500)
        with g2:  # heatmap lift
            pr = rules[(rules["Ukuran"] == 2)]
            tops = list(dict.fromkeys(list(pr.head(40)["Antecedent"]) + list(pr.head(40)["Consequent"])))[:14]
            hm = pd.DataFrame(np.nan, index=tops, columns=tops)
            for _, r in pr.iterrows():
                if r["Antecedent"] in hm.index and r["Consequent"] in hm.columns:
                    hm.loc[r["Antecedent"], r["Consequent"]] = r["Lift"]
            fig = px.imshow(hm, color_continuous_scale="Sunsetdark", aspect="auto", title="Heatmap lift (baris = antecedent, kolom = consequent)", labels=dict(color="Lift"))
            fig.update_xaxes(tickangle=-45)
            show_chart(fig, 500)
        sc1, sc2 = st.columns(2)
        with sc1:
            fig = px.scatter(rules, x="Support", y="Confidence", color="Lift", size="Lift", hover_name="Aturan", color_continuous_scale="Viridis", title="Peta aturan: Support vs Confidence (warna = lift)")
            show_chart(fig, 380)
        with sc2:
            tb = top10.iloc[::-1]
            fig = px.bar(tb, x="Lift", y="Aturan", orientation="h", text=tb["Lift"].map(lambda v: f"{v:.2f}"), title="Lift 10 aturan terbaik", color="Confidence", color_continuous_scale="Peach")
            show_chart(fig, 380)

        st.markdown("#### 6 • Interpretasi pola pembelian")
        r1_ = top10.iloc[0]
        st.markdown(
            f'<div class="ok"><b>6a. Kombinasi paling menarik:</b> <b>{r1_["Aturan"]}</b>. '
            f"<br><b>6b. Arti nilai:</b> <b>Support {r1_['Support']:.2%}</b> → {r1_['Support'] * 100:.2f}% dari seluruh nota ({angka(round(r1_['Support'] * Bm.shape[0]))} nota) memuat kedua produk sekaligus. "
            f"<b>Confidence {r1_['Confidence']:.1%}</b> → dari nota yang memuat {r1_['Antecedent']}, {r1_['Confidence']:.1%} juga memuat {r1_['Consequent']}. "
            f"<b>Lift {r1_['Lift']:.2f}</b> → peluang membeli {r1_['Consequent']} {r1_['Lift']:.2f} kali lebih besar saat {r1_['Antecedent']} dibeli dibanding peluang dasarnya (lift 1 = tidak ada hubungan).</div>", unsafe_allow_html=True)
        pairs = []
        seen = set()
        for _, r in top10.iterrows():
            key = frozenset([r["Antecedent"], r["Consequent"]])
            if key in seen:
                continue
            seen.add(key)
            pairs.append(r)
        recs = []
        for i, r in enumerate(pairs[:6], 1):
            tgt = ["Bundling paket hemat", "Cross-selling di kasir/aplikasi", "Penempatan rak berdekatan"][(i - 1) % 3]
            recs.append({"No": i, "Kombinasi produk": f"{r['Antecedent']} + {r['Consequent']}", "Strategi": tgt,
                         "Bukti (support / confidence / lift)": f"{r['Support']:.2%} / {r['Confidence']:.1%} / {r['Lift']:.2f}",
                         "Alasan": f"Dari yang membeli {r['Antecedent']}, {r['Confidence']:.0%} juga membeli {r['Consequent']}; hubungan {r['Lift']:.1f}× lebih kuat dari kebetulan."})
        st.write("")
        sec("6c–e • Rekomendasi cross-selling / bundling & penempatan rak (berdasarkan aturan)")
        show_df(pd.DataFrame(recs))
        rak = [f"**{p['Antecedent']} ↔ {p['Consequent']}** (lift {p['Lift']:.2f})" for p in pairs[:5]]
        st.markdown("**6d. Produk yang layak ditempatkan berdekatan:** " + "; ".join(rak) + ".")
        qa("Mengapa lift lebih penting daripada confidence saja?",
           "Confidence tinggi bisa menyesatkan bila consequent memang produk yang populer — banyak orang membelinya apa pun antecedent-nya. **Lift** membandingkan confidence dengan peluang dasar consequent, "
           "sehingga hanya hubungan yang benar-benar lebih kuat dari kebetulan yang diprioritaskan (lift > 1 = positif; = 1 independen; < 1 saling menjauhi).")

# ============================== TAB 5 =========================================
with tabs[4]:
    st.markdown("### 👥 Tahap 4 • Segmentasi Pelanggan — K-Means (Soal 7 & 8)")
    sec("7a • Fitur yang digunakan")
    st.markdown('<div class="info"><b>Recency</b> (apakah pelanggan masih aktif), <b>Log Frequency</b> (seberapa sering), dan <b>Log Monetary</b> (seberapa besar belanja). Ketiganya distandardisasi (Z-score) '
                "sehingga berbobot setara. RFM dipilih karena identitas pelanggan, waktu, dan nilai transaksi tersedia lengkap dan langsung berkaitan dengan tujuan bisnis (loyalitas & churn).</div>", unsafe_allow_html=True)
    st.write("")
    sec("7b–d • Elbow Method & Silhouette Score → menentukan k")
    m1, m2 = st.columns([1.3, 1])
    with m1:
        st.radio("Cara menentukan k", ["Otomatis", "Manual"], key="k_mode", horizontal=True,
                 help=f"Otomatis = k dengan Silhouette tertinggi pada rentang k = 3–6 (k = 2 terlalu kasar untuk membedakan Loyal/Potential/Churn).")
    with m2:
        st.slider("Nilai k (mode manual)", 2, 8, value=int(k_auto), key="k_pick", disabled=(k_mode == "Otomatis"))
    c1, c2, c3 = st.columns(3)
    c1.markdown(kpi("Silhouette tertinggi (k=2–8)", f"k = {best_sil_k}", f"skor {elbow['Silhouette'].max():.3f}", "blue"), unsafe_allow_html=True)
    c2.markdown(kpi("k otomatis (k=3–6)", f"k = {k_auto}", f"skor {elbow.loc[elbow['k'] == k_auto, 'Silhouette'].iloc[0]:.3f}", "pink"), unsafe_allow_html=True)
    c3.markdown(kpi("k dipakai", f"k = {k_opt}", f"Silhouette {sil_opt:.3f} • {k_mode.lower()}", "green"), unsafe_allow_html=True)
    st.write("")
    e1, e2 = st.columns(2)
    with e1:
        fig = go.Figure()
        fig.add_trace(go.Scatter(x=elbow["k"], y=elbow["Inertia"], mode="lines+markers+text", line=dict(color=C["indigo"], width=3), marker=dict(size=10),
                                 text=[f"{v:,.0f}" for v in elbow["Inertia"]], textposition="top right", name="Inertia"))
        fig.add_trace(go.Scatter(x=[k_opt], y=[elbow.loc[elbow["k"] == k_opt, "Inertia"].iloc[0]], mode="markers", name=f"k dipakai = {k_opt}",
                                 marker=dict(size=22, color=GOLD, symbol="star", line=dict(color=NAVY, width=2))))
        fig.update_layout(title="Grafik Elbow (Inertia / WCSS vs k)", xaxis_title="Jumlah klaster (k)", yaxis_title="Inertia", legend=dict(orientation="h", y=-.2))
        fig.update_xaxes(dtick=1)
        show_chart(fig, 380)
    with e2:
        fig = go.Figure(go.Bar(x=elbow["k"], y=elbow["Silhouette"], text=[f"{v:.3f}" for v in elbow["Silhouette"]], textposition="outside",
                               marker_color=[C["green"] if k_ == k_opt else C["cyan"] for k_ in elbow["k"]]))
        fig.update_layout(title="Silhouette Score per k (hijau = k dipakai)", xaxis_title="Jumlah klaster (k)", yaxis_title="Silhouette")
        fig.update_xaxes(dtick=1)
        fig.update_yaxes(range=[0, max(.6, elbow["Silhouette"].max() * 1.2)])
        show_chart(fig, 380)
    et = elbow.copy()
    et["Catatan"] = ""
    et.loc[et["k"] == best_sil_k, "Catatan"] = "Silhouette tertinggi"
    et.loc[et["k"] == k_opt, "Catatan"] = (et.loc[et["k"] == k_opt, "Catatan"] + " ⭐ dipakai").str.strip()
    show_df(et[["k", "Inertia", "Penurunan Inertia", "Penurunan (%)", "Silhouette", "Iterasi", "Catatan"]],
            column_config={"Inertia": st.column_config.NumberColumn(format="%.2f"), "Penurunan Inertia": st.column_config.NumberColumn(format="%.2f"),
                           "Penurunan (%)": st.column_config.NumberColumn(format="%.1f"), "Silhouette": st.column_config.NumberColumn(format="%.4f")})
    n_best = int(prof["Jumlah"].min())
    st.markdown(
        f'<div class="ok"><b>Keputusan: k = {k_opt}.</b><br><b>Alasan statistik:</b> pada k = {k_opt} Silhouette = {sil_opt:.3f}'
        + (f" (k = 2 memang {elbow.loc[elbow['k'] == 2, 'Silhouette'].iloc[0]:.3f}, tetapi hanya memisahkan dua kelompok sehingga tidak cukup untuk strategi berbeda)." if best_sil_k == 2 and k_opt != 2 else ".")
        + " Inertia turun tajam hingga sekitar k = 3–4, lalu melandai (siku)."
        f"<br><b>Alasan bisnis:</b> {k_opt} klaster masih mudah dipahami dan ditindaklanjuti tim pemasaran; klaster terkecil berisi {angka(n_best)} pelanggan sehingga layak jadi target kampanye.</div>", unsafe_allow_html=True)

    st.markdown("---")
    sec("7e • Hasil K-Means & proses iterasi")
    st.caption(f"k-means++ • n_init = {n_init} • seed = {seed} • inisialisasi terbaik ke-{km['init_ke']} • {'konvergen' if km['converged'] else 'berhenti'} dalam {km['iter']} iterasi • inertia akhir {km['inertia']:,.2f}")
    it = pd.DataFrame(km["hist"]).rename(columns={"Pindah": "Titik pindah klaster", "Shift": "Pergeseran centroid"})
    i1, i2 = st.columns(2)
    with i1:
        fig = px.line(it, x="Iterasi", y="Inertia", markers=True, title="Penurunan inertia per iterasi", color_discrete_sequence=[C["magenta"]])
        fig.update_xaxes(dtick=1)
        show_chart(fig, 300)
    with i2:
        fig = px.bar(it.dropna(subset=["Titik pindah klaster"]), x="Iterasi", y="Titik pindah klaster", text="Titik pindah klaster", title="Pelanggan yang pindah klaster per iterasi", color_discrete_sequence=[C["amber"]])
        fig.update_xaxes(dtick=1)
        show_chart(fig, 300)
    with st.expander("Tabel iterasi"):
        show_df(it, column_config={"Inertia": st.column_config.NumberColumn(format="%.4f"), "Pergeseran centroid": st.column_config.NumberColumn(format="%.6f")})

    sec("8a–c • Profil, label, dan karakteristik setiap klaster")
    cols = st.columns(len(prof))
    for c, (_, r) in zip(cols, prof.iterrows()):
        c.markdown(
            f'<div class="seg" style="--sc:{SEG_COLORS[r.Segmen]}"><div class="tag">Klaster {int(r.Cluster)}</div>'
            f'<h4>{SEG_ICON[r.Basis]} {r.Segmen}</h4><div class="big">{angka(r.Jumlah)}</div>'
            f'<div class="mini">pelanggan • {r.Pct_Pelanggan:.1%} basis • {r.Pct_Revenue:.1%} revenue</div>'
            f'<div class="rfm"><div><b>{dec(r.Avg_Recency, 0)}</b><span>Recency (hari)</span></div><div><b>{dec(r.Avg_Frequency, 1)}</b><span>Frequency</span></div>'
            f'<div><b>{rp_short(r.Avg_Monetary)}</b><span>Monetary</span></div></div></div>', unsafe_allow_html=True)
    st.write("")
    show_df(prof[["Cluster", "Segmen", "Jumlah", "Avg_Recency", "Avg_Frequency", "Avg_Monetary", "Z_Recency", "Z_Log_Frequency", "Z_Log_Monetary", "Pct_Pelanggan", "Pct_Revenue", "Karakteristik"]],
            column_config={"Avg_Recency": st.column_config.NumberColumn(format="%.1f"), "Avg_Frequency": st.column_config.NumberColumn(format="%.2f"), "Avg_Monetary": st.column_config.NumberColumn(format="%d"),
                           "Z_Recency": st.column_config.NumberColumn(format="%.2f"), "Z_Log_Frequency": st.column_config.NumberColumn(format="%.2f"), "Z_Log_Monetary": st.column_config.NumberColumn(format="%.2f"),
                           "Pct_Pelanggan": st.column_config.NumberColumn(format="percent"), "Pct_Revenue": st.column_config.NumberColumn(format="percent")})
    st.caption("Label ditetapkan dari bukti pusat klaster (Z-score): skor = Z_F + Z_M − Z_R tertinggi → Loyal; Recency tertinggi/F-M rendah → Churn Risk; Recency rendah tetapi F-M di bawah rata-rata → New; sisanya → Potential.")

    sec("7f • Visualisasi klaster")
    v1, v2 = st.columns(2)
    rf = rfm.copy()
    rf["Monetary (juta)"] = rf["Monetary"] / 1e6
    with v1:
        fig = px.scatter_3d(rf, x="Z_Recency", y="Z_Log_Frequency", z="Z_Log_Monetary", color="Segmen", color_discrete_map=SEG_COLORS, opacity=.8,
                            hover_data=["Customer_ID", "Recency", "Frequency"], title="Klaster pada ruang RFM terstandardisasi (3D)")
        fig.add_trace(go.Scatter3d(x=centers[:, 0], y=centers[:, 1], z=centers[:, 2], mode="markers", name="Centroid", marker=dict(size=7, color="black", symbol="diamond")))
        fig.update_layout(scene=dict(xaxis_title="Z Recency", yaxis_title="Z Log Freq", zaxis_title="Z Log Mon"))
        show_chart(fig, 500)
    with v2:
        fig = px.scatter(rf, x="Recency", y="Monetary (juta)", color="Segmen", size="Frequency", size_max=14, color_discrete_map=SEG_COLORS, opacity=.7,
                         hover_data=["Customer_ID", "Frequency"], title="Recency vs Monetary (ukuran = Frequency)")
        show_chart(fig, 500)
    pz = pd.DataFrame(centers, columns=["Recency", "Frequency (log)", "Monetary (log)"])
    pz["Segmen"] = SEG_LIST
    pz = pz.melt(id_vars="Segmen", var_name="Fitur", value_name="Z-score")
    fig = px.bar(pz, x="Fitur", y="Z-score", color="Segmen", barmode="group", color_discrete_map=SEG_COLORS, text_auto=".2f", title="Perbandingan pusat klaster (Z-score)")
    fig.add_hline(y=0, line_color="#9AA7B8")
    show_chart(fig, 360)

    sec("8d • Strategi pemasaran per segmen")
    sc = st.columns(len(prof))
    for c, (_, r) in zip(sc, prof.iterrows()):
        jud, aksi = STRATEGI[r.Basis]
        c.markdown(f'<div class="seg" style="--sc:{SEG_COLORS[r.Segmen]}"><div class="tag">{angka(r.Jumlah)} pelanggan • {r.Pct_Revenue:.0%} revenue</div>'
                   f'<h4>{SEG_ICON[r.Basis]} {r.Segmen}</h4><div class="mini"><b>{jud}</b></div><ul>{"".join(f"<li>{a}</li>" for a in aksi)}</ul></div>', unsafe_allow_html=True)
    st.write("")
    sec("8e • Segmen prioritas")
    churn_p = prof[prof["Basis"] == "Churn"]
    loyal_p = prof[prof["Basis"] == "Loyal"]
    txt = ""
    if not loyal_p.empty:
        lp = loyal_p.iloc[0]
        txt += (f"**1. {lp.Segmen} — prioritas pertahankan.** Hanya {lp.Pct_Pelanggan:.1%} pelanggan tetapi menyumbang **{lp.Pct_Revenue:.1%} revenue** "
                f"(rata-rata {rp_short(lp.Avg_Monetary)} per pelanggan). Kehilangan mereka berdampak paling besar.\n\n")
    if not churn_p.empty:
        cp = churn_p.iloc[0]
        txt += (f"**2. {cp.Segmen} — prioritas penyelamatan.** {angka(cp.Jumlah)} pelanggan ({cp.Pct_Pelanggan:.1%}) sudah rata-rata **{dec(cp.Avg_Recency, 0)} hari** tidak bertransaksi "
                f"dengan belanja historis {rp_short(cp.Avg_Monetary)}/pelanggan; semakin lama dibiarkan, peluang kembali semakin kecil.\n\n")
    oth = prof[~prof["Basis"].isin(["Loyal", "Churn"])]
    for _, r in oth.iterrows():
        txt += f"**• {r.Segmen}:** {angka(r.Jumlah)} pelanggan ({r.Pct_Pelanggan:.1%}), potensi naik kelas — prioritas menengah dengan biaya rendah (voucher & pengingat otomatis).\n\n"
    st.markdown(txt)

# ============================== TAB 6 (DASHBOARD) =============================
with tabs[5]:
    st.markdown("### 📊 Tahap 5 • Dashboard Business Intelligence (Soal 9)")
    with st.container(border=True):
        st.markdown("**🎛️ Filter interaktif**")
        f1, f2, f3, f4 = st.columns([1.2, 1.2, 1.2, 1])
        dmin, dmax = df["Tanggal"].min().date(), df["Tanggal"].max().date()
        per = f1.date_input("Periode transaksi", (dmin, dmax), min_value=dmin, max_value=dmax, key="f_per")
        cats = sorted(df["Category"].unique())
        pil_cat = f2.multiselect("Kategori produk", cats, default=cats, key="f_cat")
        prods_all = sorted(df[df["Category"].isin(pil_cat)]["Product_Name"].unique())
        pil_prod = f3.multiselect("Produk (kosong = semua)", prods_all, default=[], key="f_prod")
        pil_seg = f4.multiselect("Segmen pelanggan", SEG_LIST, default=SEG_LIST, key="f_seg")
    if isinstance(per, (tuple, list)) and len(per) == 2:
        d_from, d_to = pd.Timestamp(per[0]), pd.Timestamp(per[1])
    else:
        d_from, d_to = pd.Timestamp(dmin), pd.Timestamp(dmax)
    dd = df[(df["Tanggal"] >= d_from) & (df["Tanggal"] <= d_to) & df["Category"].isin(pil_cat) & df["Segmen"].isin(pil_seg)]
    if pil_prod:
        dd = dd[dd["Product_Name"].isin(pil_prod)]
    if dd.empty:
        st.warning("Tidak ada data untuk kombinasi filter ini.")
    else:
        txd = dd.groupby("Transaction_ID").agg(basket=("Product_ID", "nunique"), tot=("Line_Total", "sum")).reset_index()
        k1, k2, k3, k4 = st.columns(4)
        k1.markdown(kpi("💰 Total pendapatan", rp_short(dd["Line_Total"].sum()), rupiah(dd["Line_Total"].sum()), "blue"), unsafe_allow_html=True)
        k2.markdown(kpi("🧾 Jumlah transaksi", angka(len(txd)), f"rata-rata {rp_short(txd['tot'].mean())} / nota", "pink"), unsafe_allow_html=True)
        k3.markdown(kpi("👥 Pelanggan unik aktif", angka(dd["Customer_ID"].nunique()), f"{dd['Customer_ID'].nunique() / n_cust:.0%} dari basis aktif", "green"), unsafe_allow_html=True)
        k4.markdown(kpi("🛍️ Rata-rata basket_size", dec(txd["basket"].mean(), 2), "produk berbeda / nota", "orange"), unsafe_allow_html=True)
        st.write("")
        sec("📈 Analisis Penjualan")
        a1, a2 = st.columns(2)
        with a1:
            mo = dd.groupby("Bulan").agg(Revenue=("Line_Total", "sum"), Transaksi=("Transaction_ID", "nunique")).reset_index()
            fig = make_subplots(specs=[[{"secondary_y": True}]])
            fig.add_trace(go.Bar(x=mo["Bulan"], y=mo["Transaksi"], name="Transaksi", marker_color=C["cyan"], opacity=.75), secondary_y=False)
            fig.add_trace(go.Scatter(x=mo["Bulan"], y=mo["Revenue"], name="Revenue", mode="lines+markers", line=dict(color=C["magenta"], width=3), marker=dict(size=9)), secondary_y=True)
            fig.update_layout(title="Tren penjualan bulanan", legend=dict(orientation="h", y=-.15))
            fig.update_yaxes(title_text="Transaksi", secondary_y=False)
            fig.update_yaxes(title_text="Revenue (Rp)", secondary_y=True)
            show_chart(fig, 370)
        with a2:
            tp = dd.groupby("Product_Name")["Line_Total"].sum().sort_values(ascending=False).head(10).iloc[::-1].reset_index()
            fig = px.bar(tp, x="Line_Total", y="Product_Name", orientation="h", text=tp["Line_Total"].map(rp_short), title="10 produk dengan penjualan tertinggi", color="Line_Total", color_continuous_scale="Tealgrn")
            fig.update_layout(coloraxis_showscale=False, xaxis_title="Revenue (Rp)", yaxis_title="")
            show_chart(fig, 370)
        a3, a4, a5 = st.columns(3)
        with a3:
            ct = dd.groupby("Category")["Line_Total"].sum().reset_index()
            fig = px.pie(ct, names="Category", values="Line_Total", hole=.5, title="Distribusi penjualan per kategori", color_discrete_sequence=CAT_COLORS)
            fig.update_traces(textinfo="label+percent")
            show_chart(fig, 340)
        with a4:
            rg = dd.groupby("Region")["Line_Total"].sum().sort_values().reset_index()
            fig = px.bar(rg, x="Line_Total", y="Region", orientation="h", title="Penjualan per wilayah", color_discrete_sequence=[C["orange"]], text=rg["Line_Total"].map(rp_short))
            fig.update_layout(xaxis_title="Revenue (Rp)", yaxis_title="")
            show_chart(fig, 340)
        with a5:
            cm = dd.groupby("Customer_Type").agg(Revenue=("Line_Total", "sum")).reset_index()
            fig = px.pie(cm, names="Customer_Type", values="Revenue", hole=.5, title="Member vs Non-Member", color_discrete_sequence=[C["purple"], C["lime"]])
            show_chart(fig, 340)
        a6, a7 = st.columns(2)
        with a6:
            py_ = dd[dd["Payment_Method"] != UNKNOWN].groupby("Payment_Method")["Line_Total"].sum().reset_index()
            fig = px.bar(py_, x="Payment_Method", y="Line_Total", title="Penjualan per metode pembayaran", color="Payment_Method", color_discrete_sequence=[C["red"], C["sky"], C["green"], C["amber"]])
            fig.update_layout(showlegend=False, yaxis_title="Revenue (Rp)", xaxis_title="")
            show_chart(fig, 330)
        with a7:
            hm = dd.groupby(["Hari", "Jam"])["Transaction_ID"].nunique().reset_index()
            hari_ord = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]
            pv = hm.pivot(index="Hari", columns="Jam", values="Transaction_ID").reindex(hari_ord)
            fig = px.imshow(pv, aspect="auto", color_continuous_scale="YlOrRd", title="Jam & hari tersibuk (jumlah transaksi)", labels=dict(color="Nota"))
            show_chart(fig, 330)

        st.markdown("---")
        sec("🛒 Analisis Apriori (mengikuti filter di atas)")
        tx_ids = dd["Transaction_ID"].unique()
        Bd = Bm.loc[Bm.index.isin(tx_ids)]
        if pil_prod or set(pil_cat) != set(cats):
            Bd = basket_matrix(dd)
        if len(Bd) < 50 or Bd.shape[1] < 2:
            st.info("Data terfilter terlalu sedikit untuk Apriori. Perlebar filter.")
        else:
            fid, rld = apriori_run(Bd.values, tuple(Bd.columns), min_sup, min_conf, 3)
            tr = top_rules(rld, 10)
            if tr.empty:
                st.info(f"Tidak ada aturan dengan support ≥ {min_sup:.1%}, confidence ≥ {min_conf:.0%}, lift > 1 pada filter ini. Parameter diatur di Tab 4.")
            else:
                b1, b2 = st.columns([1.2, 1])
                with b1:
                    show_df(tr[["Aturan", "Support", "Confidence", "Lift"]], column_config={"Support": st.column_config.NumberColumn(format="%.4f"), "Confidence": st.column_config.NumberColumn(format="%.3f"), "Lift": st.column_config.NumberColumn(format="%.3f")})
                with b2:
                    fig = px.scatter(tr, x="Support", y="Confidence", size="Lift", color="Lift", hover_name="Aturan", color_continuous_scale="Plasma", title="Support vs Confidence (ukuran/warna = lift)")
                    show_chart(fig, 330)
                tb = tr.iloc[::-1]
                fig = px.bar(tb, x="Lift", y="Aturan", orientation="h", color="Support", color_continuous_scale="Magma_r", text=tb["Lift"].map(lambda v: f"{v:.2f}"), title="Kombinasi produk teratas berdasarkan lift")
                show_chart(fig, 400)

        st.markdown("---")
        sec("👥 Analisis K-Means")
        rr = rfm[rfm["Segmen"].isin(pil_seg) & rfm["Customer_ID"].isin(dd["Customer_ID"].unique())]
        g1, g2 = st.columns(2)
        with g1:
            cc = rr["Segmen"].value_counts().reindex(SEG_LIST).dropna().reset_index()
            cc.columns = ["Segmen", "Pelanggan"]
            fig = px.pie(cc, names="Segmen", values="Pelanggan", hole=.55, color="Segmen", color_discrete_map=SEG_COLORS, title="Jumlah pelanggan per klaster")
            fig.update_traces(textinfo="label+percent+value", sort=False)
            show_chart(fig, 360)
        with g2:
            pf = rr.groupby("Segmen").agg(Recency=("Recency", "mean"), Frequency=("Frequency", "mean"), Monetary=("Monetary", "mean")).reindex(SEG_LIST).dropna().reset_index()
            fig = make_subplots(rows=1, cols=3, subplot_titles=("Avg Recency (hari)", "Avg Frequency", "Avg Monetary (juta Rp)"))
            for j, (cn, dv) in enumerate([("Recency", 1), ("Frequency", 1), ("Monetary", 1e6)], 1):
                fig.add_trace(go.Bar(x=pf["Segmen"], y=pf[cn] / dv, marker_color=[SEG_COLORS[s] for s in pf["Segmen"]], text=[f"{v / dv:,.1f}" for v in pf[cn]], textposition="outside", showlegend=False), row=1, col=j)
            fig.update_layout(title="Karakteristik tiap klaster")
            fig.update_xaxes(showticklabels=False)
            show_chart(fig, 360)
        g3, g4 = st.columns(2)
        with g3:
            rr2 = rr.copy()
            rr2["Monetary (juta)"] = rr2["Monetary"] / 1e6
            fig = px.scatter(rr2, x="Frequency", y="Monetary (juta)", color="Segmen", color_discrete_map=SEG_COLORS, opacity=.7, hover_data=["Customer_ID", "Recency"], title="Visualisasi segmentasi: Frequency vs Monetary")
            show_chart(fig, 380)
        with g4:
            sm = dd.groupby(["Bulan", "Segmen"])["Line_Total"].sum().reset_index()
            fig = px.area(sm, x="Bulan", y="Line_Total", color="Segmen", color_discrete_map=SEG_COLORS, title="Revenue bulanan per segmen")
            fig.update_layout(yaxis_title="Revenue (Rp)")
            show_chart(fig, 380)

        st.markdown("---")
        sec("⭐ Bonus • Integrasi Apriori × K-Means (kombinasi produk relevan per segmen)")
        seg_pick = st.selectbox("Pilih segmen untuk Apriori khusus segmen", pil_seg or SEG_LIST, key="bonus_seg")
        h1, h2 = st.columns([1, 1.2])
        with h1:
            sub = dd[dd["Segmen"] == seg_pick]
            Bs = basket_matrix(sub) if len(sub) else pd.DataFrame()
            if len(Bs) >= 50:
                _, rs = apriori_run(Bs.values, tuple(Bs.columns), min_sup, min_conf, 3)
                ts = top_rules(rs, 8)
                st.markdown(f"**Aturan terbaik segmen {seg_pick}** ({angka(len(Bs))} nota)")
                if ts.empty:
                    st.info("Tidak ada aturan pada parameter saat ini. Turunkan support di Tab 4.")
                else:
                    show_df(ts[["Aturan", "Support", "Confidence", "Lift"]], column_config={"Support": st.column_config.NumberColumn(format="%.4f"), "Confidence": st.column_config.NumberColumn(format="%.3f"), "Lift": st.column_config.NumberColumn(format="%.3f")})
            else:
                st.info("Nota pada segmen ini terlalu sedikit.")
        with h2:
            pref = dd.groupby(["Segmen", "Product_Name"])["Transaction_ID"].nunique().reset_index(name="n")
            segtx = dd.groupby("Segmen")["Transaction_ID"].nunique()
            alltx = dd["Transaction_ID"].nunique()
            allp = dd.groupby("Product_Name")["Transaction_ID"].nunique() / alltx
            pref["Indeks"] = pref.apply(lambda r: (r["n"] / segtx[r["Segmen"]]) / allp[r["Product_Name"]], axis=1)
            pv = pref.pivot(index="Product_Name", columns="Segmen", values="Indeks").reindex(columns=[s for s in SEG_LIST if s in set(pref["Segmen"])])
            fig = px.imshow(pv, aspect="auto", color_continuous_scale="RdYlGn", color_continuous_midpoint=1, title="Indeks preferensi produk per segmen (1 = rata-rata; >1 = lebih sering dibeli)", labels=dict(color="Indeks"))
            show_chart(fig, 520)
        st.caption("Indeks preferensi = (% nota segmen yang memuat produk) ÷ (% semua nota yang memuat produk). Nilai > 1 berarti produk lebih relevan bagi segmen tersebut.")

        sec("🔎 Cara membaca dashboard (informasi utama untuk manajemen)")
        mo_best = mo.loc[mo["Revenue"].idxmax()]
        tp_best = dd.groupby("Product_Name")["Line_Total"].sum().idxmax()
        ct_best = dd.groupby("Category")["Line_Total"].sum().idxmax()
        st.markdown(f"- **KPI:** pada filter ini pendapatan {rupiah(dd['Line_Total'].sum())} dari {angka(len(txd))} nota dan {angka(dd['Customer_ID'].nunique())} pelanggan; basket rata-rata {dec(txd['basket'].mean(), 2)} produk.\n"
                    f"- **Penjualan:** bulan tertinggi **{mo_best['Bulan']}** ({rp_short(mo_best['Revenue'])}); produk teratas **{tp_best}**; kategori terbesar **{ct_best}**.\n"
                    "- **Apriori:** daftar aturan dan lift menunjukkan pasangan produk yang cocok di-bundling/ditata berdekatan.\n"
                    "- **K-Means:** komposisi & profil segmen menunjukkan kelompok yang perlu dipertahankan atau diselamatkan.\n"
                    "- **Bonus:** indeks preferensi menunjukkan produk apa yang layak ditawarkan ke tiap segmen.")

# ============================== TAB 7 =========================================
with tabs[6]:
    st.markdown("### 🎯 Tahap 6 • Rekomendasi Keputusan Bisnis (Soal 10)")
    st.markdown('<div class="info"><b>Pembeda penting:</b> 🟦 <b>HASIL DATA</b> = angka hasil perhitungan dari dataset • 🟨 <b>ESTIMASI</b> = proyeksi berdasar asumsi yang dapat diubah • 🟩 <b>AKTUAL</b> = baru diketahui setelah implementasi dan pengukuran.</div>', unsafe_allow_html=True)
    st.write("")
    sec("a. Strategi penempatan produk & bundling 🟦 (hasil Apriori)")
    if top10.empty:
        st.info("Belum ada aturan Apriori. Atur parameter di Tab 4.")
        pairs_u = []
    else:
        pairs_u, seen2 = [], set()
        for _, r in top10.iterrows():
            kk = frozenset([r["Antecedent"], r["Consequent"]])
            if kk not in seen2:
                seen2.add(kk)
                pairs_u.append(r)
        rows = [{"Pasangan produk": f"{p['Antecedent']} + {p['Consequent']}", "Support": p["Support"], "Confidence": p["Confidence"], "Lift": p["Lift"],
                 "Tindakan": "Bundling paket hemat + rak berdekatan" if i < 3 else "Letakkan berdekatan / tawarkan di kasir"} for i, p in enumerate(pairs_u[:6])]
        show_df(pd.DataFrame(rows), column_config={"Support": st.column_config.NumberColumn(format="%.4f"), "Confidence": st.column_config.NumberColumn(format="%.3f"), "Lift": st.column_config.NumberColumn(format="%.3f")})
    sec("b. Strategi retensi pelanggan 🟦 (hasil K-Means)")
    for _, r in prof.iterrows():
        jud, aksi = STRATEGI[r.Basis]
        st.markdown(f"**{SEG_ICON[r.Basis]} {r.Segmen}** ({angka(r.Jumlah)} pelanggan, {r.Pct_Revenue:.1%} revenue) — *{jud}*: " + " ".join(aksi[:2]))

    sec("c. Prioritas tindakan (urgensi, manfaat, biaya, kemudahan)")
    st.caption("Skor 1–5 adalah penilaian analis berdasarkan temuan data (bukan hasil perhitungan). Biaya & kemudahan: 5 = murah/mudah. Skor total = 0,3×Urgensi + 0,3×Manfaat + 0,2×Biaya + 0,2×Kemudahan.")
    churn_n = int(prof.loc[prof["Basis"] == "Churn", "Jumlah"].sum())
    loyal_pct = float(prof.loc[prof["Basis"] == "Loyal", "Pct_Revenue"].sum())
    top_pair = f"{pairs_u[0]['Antecedent']} + {pairs_u[0]['Consequent']}" if pairs_u else "pasangan produk teratas"
    pr_rows = [
        ("Win-back pelanggan Churn Risk", f"{angka(churn_n)} pelanggan lama tidak bertransaksi", 5, 4, 4, 4),
        ("Program loyalitas pelanggan Loyal", f"Loyal menyumbang {loyal_pct:.0%} revenue", 4, 5, 3, 4),
        (f"Bundling {top_pair}", "Lift tertinggi hasil Apriori", 3, 4, 4, 5),
        ("Penataan rak produk berasosiasi", "Tanpa biaya iklan, memakai aturan Apriori", 3, 3, 5, 4),
        ("Aktivasi segmen Potential/New (voucher pembelian ke-2)", "Menaikkan Frequency", 3, 4, 4, 4),
    ]
    prd = pd.DataFrame(pr_rows, columns=["Rekomendasi", "Dasar data", "Urgensi", "Manfaat", "Biaya (5=murah)", "Kemudahan"])
    prd["Skor total"] = .3 * prd["Urgensi"] + .3 * prd["Manfaat"] + .2 * prd["Biaya (5=murah)"] + .2 * prd["Kemudahan"]
    prd = prd.sort_values("Skor total", ascending=False).reset_index(drop=True)
    prd.insert(0, "Prioritas", range(1, len(prd) + 1))
    p_a, p_b = st.columns([1.5, 1])
    with p_a:
        show_df(prd, column_config={"Skor total": st.column_config.NumberColumn(format="%.2f")})
    with p_b:
        fig = px.bar(prd.iloc[::-1], x="Skor total", y="Rekomendasi", orientation="h", text=prd.iloc[::-1]["Skor total"].map(lambda v: f"{v:.2f}"), title="Skor prioritas", color="Skor total", color_continuous_scale="Sunsetdark")
        fig.update_layout(coloraxis_showscale=False, yaxis_title="")
        show_chart(fig, 330)

    sec("d. Estimasi dampak bisnis 🟨 (asumsi dapat diubah)")
    e1, e2, e3, e4 = st.columns(4)
    acc = e1.slider("Bundling: % pembeli antecedent yang menerima tawaran", 1, 30, 10) / 100
    react = e2.slider("Win-back: % Churn Risk yang kembali belanja", 1, 40, 10) / 100
    upl = e3.slider("Loyalty: kenaikan belanja pelanggan Loyal (%)", 1, 20, 5) / 100
    cost_c = e4.number_input("Biaya kampanye per pelanggan (Rp)", 1000, 100000, 10000, step=1000)
    # bundling
    inc_b, n_b = 0.0, 0
    if pairs_u:
        price = prod.set_index("Product_Name")["Reference_Unit_Price_Rp"]
        for p in pairs_u[:3]:
            ants = [a.strip() for a in p["Antecedent"].split("+")]
            cons = [c.strip() for c in p["Consequent"].split("+")]
            has_a = Bm[ants].all(axis=1)
            has_c = Bm[cons].all(axis=1)
            n_t = int((has_a & ~has_c).sum())
            n_b += n_t
            inc_b += n_t * acc * float(sum(price.get(c, 0) for c in cons))
    ch = prof[prof["Basis"] == "Churn"]
    inc_w = float((ch["Jumlah"] * ch["Avg_Monetary"]).sum()) * react
    cost_w = float(ch["Jumlah"].sum()) * cost_c
    lo = prof[prof["Basis"] == "Loyal"]
    inc_l = float(lo["Total_Revenue"].sum()) * upl
    cost_l = float(lo["Jumlah"].sum()) * cost_c
    est = pd.DataFrame({
        "Rekomendasi": [f"Bundling 3 pasangan teratas ({angka(n_b)} nota target)", "Win-back Churn Risk", "Program loyalitas Loyal"],
        "Estimasi tambahan revenue": [inc_b, inc_w, inc_l],
        "Estimasi biaya": [0.0, cost_w, cost_l],
        "Asumsi": [f"{acc:.0%} nota yang punya antecedent tapi belum punya consequent menambah 1 consequent (harga referensi)",
                   f"{react:.0%} pelanggan Churn kembali dan belanja setara rata-rata historisnya",
                   f"Belanja pelanggan Loyal naik {upl:.0%}"],
    })
    est["Estimasi manfaat bersih"] = est["Estimasi tambahan revenue"] - est["Estimasi biaya"]
    show_df(est, column_config={c: st.column_config.NumberColumn(format="%d") for c in ["Estimasi tambahan revenue", "Estimasi biaya", "Estimasi manfaat bersih"]})
    z1, z2, z3 = st.columns(3)
    z1.markdown(kpi("Estimasi tambahan revenue", rp_short(est["Estimasi tambahan revenue"].sum()), f"{est['Estimasi tambahan revenue'].sum() / total_rev:.1%} dari revenue 2025", "green"), unsafe_allow_html=True)
    z2.markdown(kpi("Estimasi biaya", rp_short(est["Estimasi biaya"].sum()), "kampanye", "orange"), unsafe_allow_html=True)
    z3.markdown(kpi("Estimasi manfaat bersih", rp_short(est["Estimasi manfaat bersih"].sum()), "tambahan revenue − biaya (belum margin)", "blue"), unsafe_allow_html=True)
    st.markdown('<div class="warn" style="margin-top:10px"><b>Catatan:</b> angka di atas hanya <b>estimasi</b> dari asumsi (bukan hasil data). Tambahan revenue belum dikurangi HPP; dampak aktual hanya dapat dipastikan lewat uji coba (A/B test) dan pengukuran setelah implementasi.</div>', unsafe_allow_html=True)

    sec("e. Rencana implementasi (setiap tindakan dibuktikan dengan data)")
    st.markdown('<div class="info"><b>Cara membaca:</b> kolom <b>Dasar data</b> berisi angka yang dihitung langsung dari dataset (🟦). '
                'Kolom <b>Target</b> adalah sasaran yang ditetapkan dari asumsi di atas (🟨), bukan hasil data. '
                'Data historis hanya membuktikan <b>siapa/apa yang layak dijadikan sasaran</b>; efek tindakan baru terbukti setelah uji coba (🟩), karena itu setiap tindakan punya syarat lanjut/hentikan.</div>',
                unsafe_allow_html=True)
    st.write("")

    # ---------- bukti dari data ----------
    HARI_ID = {"Monday": "Senin", "Tuesday": "Selasa", "Wednesday": "Rabu", "Thursday": "Kamis",
               "Friday": "Jumat", "Saturday": "Sabtu", "Sunday": "Minggu"}
    basis_cust = rfm["Segmen"].map(base_label)
    ch_c = rfm[basis_cust == "Churn"]
    lo_c = rfm[basis_cust == "Loyal"]
    nw_c = rfm[basis_cust.isin(["New", "Potential"])]

    gap = tx.sort_values(["Customer_ID", "Tanggal"]).groupby("Customer_ID")["Tanggal"].diff().dt.days.dropna()
    med_gap = float(gap.median()) if len(gap) else float("nan")
    gap_txt = f"{dec(med_gap, 0)} hari" if not np.isnan(med_gap) else "n/a"

    jam_tx = df.groupby("Jam")["Transaction_ID"].nunique()
    hari_tx = df.groupby("Hari")["Transaction_ID"].nunique()
    bln_rev = df.groupby("Bulan_No")["Line_Total"].sum()
    jam_ramai, jam_sepi = int(jam_tx.idxmax()), int(jam_tx.idxmin())
    hari_ramai, hari_sepi = HARI_ID.get(hari_tx.idxmax(), hari_tx.idxmax()), HARI_ID.get(hari_tx.idxmin(), hari_tx.idxmin())
    bln_top, bln_low = BULAN[int(bln_rev.idxmax()) - 1], BULAN[int(bln_rev.idxmin()) - 1]

    pair_ev = []
    for p in pairs_u[:3]:
        ants = [a.strip() for a in p["Antecedent"].split("+")]
        cons = [c.strip() for c in p["Consequent"].split("+")]
        ha, hc = Bm[ants].all(axis=1), Bm[cons].all(axis=1)
        pair_ev.append({"nama": f"{p['Antecedent']} + {p['Consequent']}", "n_ant": int(ha.sum()), "n_both": int((ha & hc).sum()),
                        "n_target": int((ha & ~hc).sum()), "sup": p["Support"], "conf": p["Confidence"], "lift": p["Lift"]})
    pair1 = pair_ev[0] if pair_ev else None
    nama_pairs = "; ".join(e["nama"] for e in pair_ev) if pair_ev else "pasangan produk teratas (atur parameter di Tab 4)"
    ev_pair_txt = (" | ".join(f"{e['nama']}: lift {e['lift']:.2f}, confidence {e['conf']:.1%}, {angka(e['n_ant'])} nota berisi antecedent "
                              f"(→ {angka(e['n_target'])} nota belum membeli pasangannya)" for e in pair_ev) if pair_ev else "Belum ada aturan Apriori")
    attach_base = float(np.mean([e["conf"] for e in pair_ev])) if pair_ev else float("nan")
    attach_tgt = attach_base + (1 - attach_base) * acc if pair_ev else float("nan")
    basket_gain = (sum(e["n_target"] for e in pair_ev) * acc / len(Bm)) if pair_ev else 0.0

    ch_rev = float(ch_c["Monetary"].sum())
    ch_pool = ch_c[ch_c["Monetary"] >= ch_c["Monetary"].median()] if len(ch_c) else ch_c   # 50% churn bernilai tertinggi
    pool_share = float(ch_pool["Monetary"].sum()) / ch_rev if ch_rev > 0 else 0.0
    ch_rec = float(ch_c["Recency"].mean()) if len(ch_c) else float("nan")
    rec_x = (ch_rec / med_gap) if (len(ch_c) and med_gap and not np.isnan(med_gap)) else float("nan")
    one_timer = int((nw_c["Frequency"] == 1).sum())
    lo_share = float(lo_c["Monetary"].sum()) / total_rev if total_rev else 0.0
    lo_freq = float(lo_c["Frequency"].mean()) if len(lo_c) else float("nan")
    cost_pilot = len(ch_pool) * cost_c

    # ---------- tabel rencana ----------
    plan = [
        {"Fase": "0 • Persiapan", "Mulai": 1, "Selesai": 2, "Tindakan": "Bekukan baseline KPI, tetapkan kelompok uji & kontrol, ekspor daftar pelanggan per segmen dari sistem ini.",
         "Dasar data": f"{angka(n_cust)} pelanggan aktif sudah bersegmen; {angka(len(Bm))} nota; baseline basket_size {dec(avg_basket, 2)} produk/nota.",
         "Target": "Baseline & daftar sasaran siap sebelum tindakan apa pun dijalankan.", "Penanggung jawab": "Analis data",
         "Biaya / sumber daya": "Rp 0 (memakai ekspor Excel sistem)", "Syarat lanjut / hentikan": "Lanjut bila baseline dan daftar sasaran disetujui manajer toko."},
        {"Fase": "1 • Quick win", "Mulai": 3, "Selesai": 4, "Tindakan": f"Tata rak/etalase agar pasangan produk berdekatan: {nama_pairs}. Kerjakan pada {hari_sepi} jam {jam_sepi}:00 (transaksi paling sedikit di data) agar tidak mengganggu pelanggan.",
         "Dasar data": ev_pair_txt, "Target": (f"Attach rate pasangan naik dari {attach_base:.1%} menjadi ±{attach_tgt:.1%} (asumsi {acc:.0%} nota menerima)." if pair_ev else "-"),
         "Penanggung jawab": "Manajer toko", "Biaya / sumber daya": "Tenaga toko 1–2 hari; tanpa biaya iklan",
         "Syarat lanjut / hentikan": "Pertahankan bila attach rate naik ≥ separuh target setelah 4 minggu dan penjualan produk antecedent tidak turun; selain itu kembalikan tata letak."},
        {"Fase": "2 • Bundling", "Mulai": 5, "Selesai": 10, "Tindakan": f"Luncurkan paket hemat 3 pasangan teratas (tawarkan di kasir). Jalankan pekan ON/OFF bergantian sebagai pembanding. Pastikan sudah berjalan sebelum bulan puncak historis ({bln_top}).",
         "Dasar data": (f"{angka(sum(e['n_target'] for e in pair_ev))} nota sudah memuat antecedent tetapi belum consequent = pasar tambahan; bulan terendah {bln_low}, tertinggi {bln_top}." if pair_ev else "Belum ada aturan Apriori"),
         "Target": f"basket_size naik dari {dec(avg_basket, 2)} menjadi ±{dec(avg_basket + basket_gain, 2)}; estimasi tambahan revenue {rp_short(est['Estimasi tambahan revenue'].iloc[0])}.",
         "Penanggung jawab": "Pemasaran & merchandising", "Biaya / sumber daya": "Diskon paket (turunkan margin; hitung HPP sebelum luncur)",
         "Syarat lanjut / hentikan": "Lanjut bila basket_size pekan ON > pekan OFF dan margin paket tetap positif; hentikan paket yang attach rate-nya tidak naik."},
        {"Fase": "2 • Win-back (pilot)", "Mulai": 5, "Selesai": 10, "Tindakan": f"Kirim voucher berbatas waktu via WhatsApp ke {angka(len(ch_pool))} pelanggan Churn Risk dengan Monetary historis tertinggi. Acak 50:50 menjadi kelompok uji dan kontrol (kontrol tidak dikirimi).",
         "Dasar data": (f"{angka(len(ch_c))} pelanggan Churn Risk, rata-rata {dec(ch_rec, 0)} hari tidak belanja = ±{dec(rec_x, 1)}× median jeda belanja normal ({gap_txt}). "
                        f"Separuh teratas ({angka(len(ch_pool))} orang) mewakili {pool_share:.0%} belanja historis segmen ini ({rp_short(ch_rev)} total)."),
         "Target": f"Reaktivasi uji ≥ {react:.0%} dan lebih tinggi ≥ 5 poin persentase dari kontrol.",
         "Penanggung jawab": "Tim pemasaran", "Biaya / sumber daya": f"≤ {rp_short(cost_pilot)} ({angka(len(ch_pool))} × {rupiah(cost_c)})",
         "Syarat lanjut / hentikan": "Perluas ke seluruh Churn Risk bila reaktivasi uji − kontrol ≥ 5 poin dan ROI > 0; hentikan bila selisih < 2 poin."},
        {"Fase": "3 • Aktivasi", "Mulai": 7, "Selesai": 14, "Tindakan": f"Kirim voucher pembelian ke-2 dengan minimum belanja kepada pelanggan New/Potential, ±{dec(med_gap, 0) if not np.isnan(med_gap) else 'n/a'} hari setelah belanja terakhir (sesuai ritme belanja normal), plus rekomendasi produk dari aturan Apriori.",
         "Dasar data": f"{angka(len(nw_c))} pelanggan New/Potential, {angka(one_timer)} di antaranya baru 1× belanja; median jeda antar belanja = {gap_txt}.",
         "Target": f"Naikkan Frequency; target persentase pembeli satu-kali yang belanja lagi ditetapkan manajemen (acuan awal ≥ {react:.0%}).",
         "Penanggung jawab": "Tim pemasaran / CRM", "Biaya / sumber daya": f"Voucher hanya dibiayai bila ditebus (maks. {angka(len(nw_c))} × {rupiah(cost_c)})",
         "Syarat lanjut / hentikan": "Lanjut bila tingkat penebusan ≥ target dan Frequency kelompok uji > kontrol; hentikan voucher bila penebusan < 3%."},
        {"Fase": "3 • Loyalitas", "Mulai": 11, "Selesai": 20, "Tindakan": "Bangun program poin/member tier & akses awal promo untuk pelanggan Loyal. Jangan beri diskon besar tanpa syarat.",
         "Dasar data": f"{angka(len(lo_c))} pelanggan Loyal menyumbang {lo_share:.1%} revenue; rata-rata {dec(lo_freq, 1)} nota/pelanggan.",
         "Target": f"Monetary rata-rata Loyal naik {upl:.0%} (dari {rp_short(lo_c['Monetary'].mean()) if len(lo_c) else '-'}).",
         "Penanggung jawab": "Tim pemasaran / CRM", "Biaya / sumber daya": f"≤ {rp_short(len(lo_c) * cost_c)} ({angka(len(lo_c))} × {rupiah(cost_c)})",
         "Syarat lanjut / hentikan": "Evaluasi tiap bulan; hentikan bila biaya poin > tambahan margin dari kenaikan belanja."},
        {"Fase": "4 • Otomasi & evaluasi", "Mulai": 17, "Selesai": 24, "Tindakan": "Jalankan ulang pipeline (cleaning → RFM → K-Means → Apriori) setiap bulan dengan data terbaru, pantau migrasi segmen dan KPI.",
         "Dasar data": f"Sistem ini sudah memproses {angka(len(raw))} baris mentah secara otomatis; segmen dapat berubah seiring waktu (Recency terus bertambah).",
         "Target": "Laporan bulanan KPI & migrasi segmen (Churn → kembali, Potential → Loyal).",
         "Penanggung jawab": "Analis data", "Biaya / sumber daya": "Waktu analis ±½ hari/bulan",
         "Syarat lanjut / hentikan": "Revisi strategi bila ukuran/profil segmen berubah > 20% antar bulan."},
    ]
    plan_df = pd.DataFrame(plan)
    show_df(plan_df.drop(columns=["Mulai", "Selesai"]).assign(Minggu=[f"{p['Mulai']}–{p['Selesai']}" for p in plan])[
        ["Fase", "Minggu", "Tindakan", "Dasar data", "Target", "Penanggung jawab", "Biaya / sumber daya", "Syarat lanjut / hentikan"]],
        column_config={c: st.column_config.TextColumn(width="large") for c in ["Tindakan", "Dasar data", "Target", "Syarat lanjut / hentikan"]})

    gcols = [C["indigo"], C["orange"], C["green"], C["red"], C["purple"], C["cyan"], C["slate"]]
    gf = go.Figure()
    for i, p in enumerate(plan):
        gf.add_trace(go.Bar(y=[f"{p['Fase']}"], x=[p["Selesai"] - p["Mulai"] + 1], base=[p["Mulai"] - 1], orientation="h", marker_color=gcols[i % len(gcols)],
                            text=f"Mgg {p['Mulai']}–{p['Selesai']}", textposition="inside", showlegend=False, hovertext=p["Tindakan"], hoverinfo="text"))
    gf.update_layout(title="Jadwal implementasi (24 minggu)", barmode="overlay")
    gf.update_xaxes(title="Minggu ke-", range=[0, 24], dtick=2)
    gf.update_yaxes(autorange="reversed")
    show_chart(gf, 360)

    st.markdown("**Bukti vs batas bukti** (agar klaim tidak berlebihan)")
    show_df(pd.DataFrame({
        "Keputusan": ["Penataan rak & bundling", "Win-back Churn Risk", "Program loyalitas Loyal", "Aktivasi New/Potential"],
        "Yang dibuktikan data (🟦)": [
            f"Pasangan produk dibeli bersamaan lebih sering dari kebetulan (lift > 1); {angka(sum(e['n_target'] for e in pair_ev))} nota berpeluang ditambah." if pair_ev else "Belum ada aturan.",
            f"Segmen ini ±{dec(rec_x, 1)}× lebih lama tidak belanja dari ritme normal dan memiliki {rp_short(ch_rev)} belanja historis." if len(ch_c) else "Tidak ada klaster churn pada k ini.",
            f"Segmen kecil ({angka(len(lo_c))} orang) menyumbang {lo_share:.0%} revenue → risiko terbesar bila hilang.",
            f"{angka(one_timer)} pelanggan baru 1× belanja; belum terbentuk kebiasaan."],
        "Belum dibuktikan (🟩 perlu uji)": [
            "Apriori menunjukkan asosiasi, bukan sebab-akibat: rak/bundling belum tentu menaikkan pembelian → uji pekan ON/OFF.",
            "Belum diketahui berapa yang akan kembali karena voucher → uji kelompok uji vs kontrol.",
            "Belum diketahui apakah poin menaikkan belanja atau hanya memberi diskon pada yang tetap akan belanja → bandingkan dengan Loyal non-peserta.",
            "Belum diketahui tingkat penebusan voucher → ukur pada pilot."]}),
        column_config={c: st.column_config.TextColumn(width="large") for c in ["Yang dibuktikan data (🟦)", "Belum dibuktikan (🟩 perlu uji)"]})

    sec("f. KPI evaluasi keberhasilan")
    roi_est = (est["Estimasi manfaat bersih"].sum() / est["Estimasi biaya"].sum()) if est["Estimasi biaya"].sum() > 0 else float("nan")
    show_df(pd.DataFrame({
        "KPI": ["Rata-rata basket_size", "Attach rate pasangan bundling", "Reactivation rate (Churn → kembali belanja)", "Rata-rata Frequency (nota/pelanggan)",
                "Monetary rata-rata pelanggan Loyal", "ROI kampanye"],
        "Baseline (hasil data 🟦)": [dec(avg_basket, 2), f"{attach_base:.1%}" if pair_ev else "-", "Kelompok kontrol (tidak dikirimi)", dec(rfm["Frequency"].mean(), 2),
                                     rp_short(lo_c["Monetary"].mean()) if len(lo_c) else "-", "-"],
        "Target (asumsi 🟨)": [dec(avg_basket + basket_gain, 2), f"{attach_tgt:.1%}" if pair_ev else "-", f"≥ {react:.0%} & ≥ kontrol + 5 poin", "Naik dibanding baseline (target ditetapkan manajemen)",
                               rp_short(lo_c["Monetary"].mean() * (1 + upl)) if len(lo_c) else "-", f"> 0 (estimasi {dec(roi_est, 1)}×)" if not np.isnan(roi_est) else "> 0"],
        "Cara mengukur": ["Σ produk berbeda ÷ jumlah nota, per pekan", "% nota yang memuat antecedent dan juga consequent", "Pelanggan uji yang bertransaksi ≤ 30 hari setelah kampanye ÷ total dikirimi, dibanding kontrol",
                          "Rata-rata nota per pelanggan per periode (hasil RFM ulang bulanan)", "Σ Line_Total Loyal ÷ jumlah Loyal, bandingkan sebelum/sesudah dan dengan non-peserta", "(Tambahan revenue − biaya) ÷ biaya, setelah dikurangi HPP"],
        "Frekuensi & sumber data": ["Mingguan • data kasir", "Mingguan • data kasir", "Pada minggu ke-4 & ke-8 pilot • WhatsApp + data kasir", "Bulanan • sistem ini", "Bulanan • sistem ini", "Akhir pilot (minggu 10) & tiap kuartal"]}),
        column_config={c: st.column_config.TextColumn(width="large") for c in ["Cara mengukur", "Target (asumsi 🟨)"]})
    st.markdown("#### 📥 Unduh luaran")
    dl1, dl2 = st.columns(2)
    dq_export = cmp_df.drop(columns="Status")
    rfm_exp = rfm.drop(columns=[c for c in ["Cluster"] if False])
    with dl1:
        st.download_button("⬇️ Dataset hasil preprocessing + dokumentasi (Excel)",
                           to_excel_bytes({"Data Bersih": df.drop(columns=["Segmen", "Tanggal", "Bulan_No"]), "Dokumentasi Cleaning": clean_log, "Sebelum vs Sesudah": dq_export}),
                           "dataset_bersih_toko.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    with dl2:
        st.download_button("⬇️ Seluruh hasil analisis (RFM, Apriori, K-Means, rekomendasi)",
                           to_excel_bytes({"RFM & Segmen": rfm_exp, "Transaksi (basket_size)": tx, "Frequent Itemsets": fi, "Aturan Asosiasi": rules.drop(columns=["Ukuran"]) if not rules.empty else pd.DataFrame(),
                                           "Elbow & Silhouette": elbow, "Profil Klaster": prof.drop(columns=["Basis"]), "Prioritas": prd, "Estimasi Dampak": est}),
                           "hasil_analisis_uts_bi.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

st.markdown("---")
st.caption(f"© UTS {matkul} • {nama} ({npm}) • {kelas} / {semester} • Universitas Timor • Dibangun dengan Streamlit")
